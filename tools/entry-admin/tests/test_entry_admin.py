"""Run from tools/entry-admin:   python -m unittest discover -s tests -v

Builds the fake sample data in a temp folder and checks every flag, plus re-runs, emails and the judging pack.
"""
from __future__ import annotations

import csv
import re
import sys
import tomllib
import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook
from pypdf import PdfReader

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))
sys.path.insert(0, str(ROOT / "sample"))

import make_sample  # noqa: E402
from entry_admin import run  # noqa: E402
from entryadmin import codes, config, emails, pack, tracker  # noqa: E402
from entryadmin.config import ConfigError  # noqa: E402
from entryadmin.xlsxutil import read_table  # noqa: E402

FORBIDDEN = ("charity", "non-profit", "nonprofit", "not-for-profit", "ngo", "foundation", "society",
              "united nations", "global goals")


class Base(unittest.TestCase):
    def setUp(self):
        self._tmp = tempfile.TemporaryDirectory()
        self.dir = Path(self._tmp.name)
        self.info = make_sample.build(self.dir)
        self.n = len(self.info["entries"])
        self.cfg = config.load(self.dir / "config.toml")
        self.addCleanup(self._tmp.cleanup)

    def process(self):
        return tracker.process(self.cfg)

    def table(self):
        return read_table(self.cfg.path("tracker"))

    def by_key(self):
        """sample key (e.g. 'over') -> tracker row."""
        sid_to_key = dict(self.info["entries"])
        return {sid_to_key[r["Submission ID"]]: r for r in self.table() if r["Submission ID"] in sid_to_key}

    def flags(self):
        return {k: {key for key, _ in tracker.raised_flags(r)} for k, r in self.by_key().items()}

    def csv_rows(self):
        with (self.dir / "tally-export.csv").open(encoding="utf-8-sig", newline="") as f:
            return list(csv.DictReader(f))

    def write_csv_rows(self, rows):
        with (self.dir / "tally-export.csv").open("w", newline="", encoding="utf-8-sig") as f:
            w = csv.DictWriter(f, fieldnames=list(rows[0]))
            w.writeheader()
            w.writerows(rows)

    def edit_tracker(self, **by_key_edits):
        """by_key_edits: sample key -> {column: value}. Types into the tracker like the owner would."""
        path = self.cfg.path("tracker")
        wb = load_workbook(path)
        ws = wb.active
        headers = [c.value for c in ws[1]]
        sid_to_key = dict(self.info["entries"])
        for row in ws.iter_rows(min_row=2):
            key = sid_to_key.get(row[headers.index("Submission ID")].value)
            for col, val in by_key_edits.get(key, {}).items():
                row[headers.index(col)].value = val
        wb.save(path)


class TestFlags(Base):
    def test_every_flag(self):
        self.process()
        f = self.flags()
        expected = {
            "clean": set(), "limit800": set(), "age20": set(), "half_ok": set(), "half_spaces": set(),
            "over": {"over_word_limit"}, "limit801": {"over_word_limit"},
            "scanned": {"pdf_problem"}, "corrupt": {"pdf_problem"}, "nofile": {"pdf_problem"},
            "dupemail": {"duplicate_entrant"},
            "dupname_a": {"duplicate_entrant"}, "dupname_b": {"duplicate_entrant"},
            "old": set(),
            "half_reuse": {"fee_code"}, "half_bad": {"fee_code"}, "half_none": {"fee_code"},
            "noproof": {"no_proof"},
            "txn_a": {"duplicate_transfer"}, "txn_b": {"duplicate_transfer"},
            "multi": {"over_word_limit", "age_unclear", "no_proof"},
            "fee_other": {"fee_code"}, "nonlatin": {"name_not_latin"}, "accents": set(),
            "gdoc": {"pdf_problem"}, "other_url": {"pdf_problem"}, "few_words": {"pdf_problem"},
            "docx": {"pdf_problem"}, "nonlatin_multi": {"name_not_latin", "age_unclear"},
        }
        # 'clean' shares an email with 'dupemail', so both are flagged
        expected["clean"] = {"duplicate_entrant"}
        self.assertEqual(f, expected)

    def test_flag_details(self):
        self.process()
        r = self.by_key()
        self.assertTrue(r["scanned"]["PDF problem"].startswith("Scanned"))
        self.assertTrue(r["corrupt"]["PDF problem"].startswith("Unreadable"))
        self.assertTrue(r["nofile"]["PDF problem"].startswith("Missing"))
        self.assertTrue(r["half_reuse"]["Fee code problem"].startswith("Used"))
        self.assertIn("ZH-013", r["half_reuse"]["Fee code problem"])  # first user of the code
        self.assertTrue(r["half_bad"]["Fee code problem"].startswith("Invalid"))
        self.assertTrue(r["half_none"]["Fee code problem"].startswith("Missing"))
        self.assertTrue(r["multi"]["Age unclear"].startswith("Unclear"))
        self.assertTrue(r["fee_other"]["Fee code problem"].startswith("Unclear"))
        self.assertIn("李明", r["nonlatin"]["Non-Latin name"])
        for k in ("gdoc", "other_url", "few_words", "docx"):  # anything that isn't a PDF upload
            self.assertTrue(r[k]["PDF problem"].startswith("Not a PDF"), k)
            self.assertEqual(r[k]["Word count"], "")
        self.assertIn("docs.google.com", r["gdoc"]["PDF problem"])
        self.assertIn("see my email", r["few_words"]["PDF problem"])

    def test_word_count_is_exact_on_simple_pdfs(self):
        self.process()
        r = self.by_key()
        self.assertEqual(r["limit800"]["Word count"], 800)
        self.assertEqual(r["limit801"]["Word count"], 801)

    def test_unreadable_pdf_gets_no_word_count(self):
        self.process()
        self.assertEqual(self.by_key()["corrupt"]["Word count"], "")

    def test_code_usage_written_to_codes_file(self):
        self.process()
        used = {r["Code"]: r["Used by submission"] for r in read_table(self.cfg.path("codes"))}
        self.assertEqual(used[self.info["codes"][0]], "ZH-013")
        self.assertEqual(used[self.info["codes"][1]], "ZH-017")
        self.assertEqual(used[self.info["codes"][2]], "")

    def test_code_marked_used_by_hand_flags_new_users(self):
        wb = load_workbook(self.cfg.path("codes"))
        ws = wb.active
        for row in ws.iter_rows(min_row=2):
            if row[0].value == self.info["codes"][1]:
                row[4].value = "ZH-099"  # redeemed some other way
        wb.save(self.cfg.path("codes"))
        self.process()
        self.assertEqual(self.flags()["half_spaces"], {"fee_code"})

    def test_two_entries_pointing_at_same_file_are_flagged(self):
        src = self.dir / "tally-export.csv"
        rows = self.csv_rows()
        rows[1]["Your essay (PDF)"] = rows[0]["Your essay (PDF)"]
        self.write_csv_rows(rows)
        self.process()
        f = self.flags()
        self.assertIn("pdf_problem", f["clean"])
        self.assertIn("pdf_problem", f["over"])

    def test_essay_answer_kinds(self):
        from entryadmin.pdfs import classify_essay
        pdf = "https://storage.tally.so/private/My%20Essay.PDF?id=1&signature=x"
        self.assertEqual(classify_essay("")[0], "empty")
        self.assertEqual(classify_essay("   ")[0], "empty")
        self.assertEqual(classify_essay(pdf)[0], "upload")
        self.assertEqual(classify_essay(pdf + ", " + pdf)[0], "upload")
        for bad in ("https://docs.google.com/document/d/abc/edit", "https://example.com/essay",
                    "https://storage.tally.so/private/essay.docx?id=1", "https://storage.tally.so/private/essay.pdf.exe",
                    "here you go", " ".join(["a long pasted essay"] * 100)):
            self.assertEqual(classify_essay(bad)[0], "notpdf", bad)

    def test_links_are_never_fetched(self):
        import socket
        def boom(*a, **k):
            raise AssertionError("network used")
        real = socket.socket.connect
        socket.socket.connect = boom
        try:
            self.process()
        finally:
            socket.socket.connect = real
        self.assertEqual(self.flags()["gdoc"], {"pdf_problem"})

    def test_pasted_text_is_a_pdf_problem_with_no_word_count(self):
        rows = self.csv_rows()
        rows[0]["Your essay (PDF)"] = " ".join(["word"] * 850)
        self.write_csv_rows(rows)
        self.process()
        r = self.by_key()["clean"]
        self.assertTrue(r["PDF problem"].startswith("Not a PDF"))
        self.assertEqual(r["Word count"], "")
        self.assertEqual(r["Over word limit"], "")

    def test_blank_essay_is_a_pdf_problem(self):
        rows = self.csv_rows()
        rows[0]["Your essay (PDF)"] = ""
        self.write_csv_rows(rows)
        self.process()
        self.assertTrue(self.by_key()["clean"]["PDF problem"].startswith("Missing"))

    def test_non_latin_names(self):
        from entryadmin.tracker import non_latin_chars
        self.assertEqual(non_latin_chars("Ada Lovelace-Smith O'Neil"), "")
        self.assertEqual(non_latin_chars("José Álvarez"), "")
        self.assertEqual(non_latin_chars("Анна Ivanova"), "Анна")
        self.assertEqual(non_latin_chars("محمد Ali"), "محمد")
        self.assertEqual(non_latin_chars("김민수"), "김민수")

    def test_fee_answers_match_exactly_but_ignore_case_and_spacing(self):
        rows = self.csv_rows()
        rows[0]["Your entry fee"] = "  rm 40 (ABOUT usd 10) "
        rows[1]["Your entry fee"] = "Half fee"            # not the exact option text
        self.write_csv_rows(rows)
        self.process()
        f = self.flags()
        self.assertNotIn("fee_code", f["clean"])
        self.assertIn("fee_code", f["over"])

    def test_missing_codes_file_is_reported(self):
        self.cfg.path("codes").unlink()
        r = self.process()
        self.assertTrue(any("No codes file" in n for n in r["notes"]))
        self.assertEqual(self.flags()["half_ok"], {"fee_code"})


class TestColumnMapping(Base):
    def test_renamed_question_is_a_one_line_fix(self):
        src = self.dir / "tally-export.csv"
        src.write_text(src.read_text(encoding="utf-8-sig").replace("Full name", "Your full name"), encoding="utf-8")
        with self.assertRaises(ConfigError) as cm:
            self.process()
        self.assertIn("full_name", str(cm.exception))
        self.assertIn("Your full name", str(cm.exception))  # shows what the CSV actually has
        cfg_file = self.dir / "config.toml"
        cfg_file.write_text(cfg_file.read_text().replace('full_name       = "Full name"',
                                                         'full_name       = "Your full name"'))
        self.cfg = config.load(cfg_file)
        self.process()
        self.assertEqual(self.by_key()["clean"]["Full name"], "Ada Example")

    def test_header_matching_ignores_case_and_extra_spaces(self):
        src = self.dir / "tally-export.csv"
        src.write_text(src.read_text(encoding="utf-8-sig").replace("Full name", "FULL  NAME"), encoding="utf-8")
        self.process()
        self.assertEqual(self.by_key()["clean"]["Full name"], "Ada Example")

    def test_semicolon_csv_from_excel(self):
        src = self.dir / "tally-export.csv"
        with src.open(encoding="utf-8-sig", newline="") as f:
            rows = list(csv.reader(f))
        with src.open("w", newline="", encoding="utf-8-sig") as f:
            csv.writer(f, delimiter=";").writerows(rows)
        self.process()
        self.assertEqual(len(self.table()), self.n)


class TestRerun(Base):
    def test_owner_columns_survive_and_only_new_entries_are_added(self):
        self.process()
        before = {r["Submission ID"]: r["Entry ID"] for r in self.table()}
        self.edit_tracker(clean={"Payment verified (Y/N)": "Y", "Notes": "Checked bank transfer", "Status": "Confirmed"},
                          over={"Notes": "Asked to shorten", "Status": "Pending"})
        # a new entry arrives, submitted LATER than the rest, and another one EARLIER
        src = self.dir / "tally-export.csv"
        rows = self.csv_rows()
        new = dict(rows[0], **{"Submission ID": "subNEW", "Submitted at": "2026-12-30 10:00:00",
                               "Full name": "New Person", "Email": "new@example.com", "Transfer reference": "TXNNEW"})
        early = dict(rows[0], **{"Submission ID": "subEARLY", "Submitted at": "2026-01-01 10:00:00",
                                 "Full name": "Early Person", "Email": "early@example.com", "Transfer reference": "TXNEARLY"})
        self.write_csv_rows([early] + rows + [new])
        r = self.process()
        self.assertEqual(sorted(r["new"]), [f"ZH-{self.n + 1:03d}", f"ZH-{self.n + 2:03d}"])
        after = {x["Submission ID"]: x for x in self.table()}
        for sid, eid in before.items():  # nobody renumbered
            self.assertEqual(after[sid]["Entry ID"], eid)
        self.assertEqual(len(after), self.n + 2)
        k = {v: kk for kk, v in dict(self.info["entries"]).items()}
        self.assertEqual(after[k["clean"]]["Payment verified (Y/N)"], "Y")
        self.assertEqual(after[k["clean"]]["Notes"], "Checked bank transfer")
        self.assertEqual(after[k["clean"]]["Status"], "Confirmed")
        self.assertEqual(after[k["over"]]["Notes"], "Asked to shorten")

    def test_rerun_with_no_changes_is_stable(self):
        self.process()
        a = self.table()
        r = self.process()
        self.assertEqual(r["new"], [])
        self.assertEqual(a, self.table())

    def test_new_duplicate_flags_older_entry_too(self):
        self.process()
        self.assertNotIn("duplicate_entrant", self.flags()["old"])
        src = self.dir / "tally-export.csv"
        rows = self.csv_rows()
        rows.append(dict(rows[0], **{"Submission ID": "subDUP", "Full name": "Someone Else",
                                     "Email": "JON@example.com ", "Transfer reference": "TXNDUP"}))
        self.write_csv_rows(rows)
        self.process()
        self.assertIn("duplicate_entrant", self.flags()["old"])

    def test_entry_missing_from_csv_is_kept(self):
        self.process()
        src = self.dir / "tally-export.csv"
        lines = src.read_text(encoding="utf-8-sig").splitlines()
        src.write_text("\n".join(lines[:-1]), encoding="utf-8")
        r = self.process()
        self.assertEqual(len(self.table()), self.n)
        self.assertTrue(any("not in this CSV" in n for n in r["notes"]))

    def test_extra_columns_added_by_owner_are_kept(self):
        self.process()
        path = self.cfg.path("tracker")
        wb = load_workbook(path)
        ws = wb.active
        ws.cell(row=1, column=ws.max_column + 1, value="Judge assigned")
        ws.cell(row=2, column=ws.max_column, value="Judge 3")
        wb.save(path)
        self.process()
        self.assertEqual(self.table()[0]["Judge assigned"], "Judge 3")

    def test_colour_coding(self):
        self.process()
        ws = load_workbook(self.cfg.path("tracker")).active
        headers = [c.value for c in ws[1]]
        sid_to_row = {ws.cell(row=i, column=2).value: i for i in range(2, ws.max_row + 1)}
        col = headers.index("Over word limit") + 1
        flagged = ws.cell(row=sid_to_row["sub002"], column=col)
        clean = ws.cell(row=sid_to_row["sub001"], column=col)
        self.assertEqual(flagged.fill.start_color.rgb[-6:], "F8CBAD")
        self.assertNotEqual(clean.fill.fill_type, "solid")
        dup = ws.cell(row=sid_to_row["sub008"], column=headers.index("Duplicate entrant") + 1)
        self.assertEqual(dup.fill.start_color.rgb[-6:], "FFE699")


class TestCodes(Base):
    def test_generate_unique_readable_and_appended(self):
        path = self.dir / "fresh-codes.xlsx"
        first = codes.generate(path, 500)
        more = codes.generate(path, 500)
        everything = first + more
        self.assertEqual(len(set(everything)), 1000)
        for c in everything:
            self.assertRegex(c, r"^[A-HJKMNP-Z2-9]{4}-[A-HJKMNP-Z2-9]{4}$")
            for bad in "0O1IL":
                self.assertNotIn(bad, c)
        rows = read_table(path)
        self.assertEqual(len(rows), 1000)
        self.assertEqual(list(rows[0]), ["Code", "Issued to", "Email", "Date issued", "Used by submission"])

    def test_generate_keeps_owner_typing(self):
        path = self.dir / "fresh-codes.xlsx"
        codes.generate(path, 3)
        wb = load_workbook(path)
        wb.active["B2"] = "Mr Example"
        wb.active["C2"] = "x@example.com"
        wb.save(path)
        codes.generate(path, 2)
        rows = read_table(path)
        self.assertEqual((rows[0]["Issued to"], rows[0]["Email"]), ("Mr Example", "x@example.com"))
        self.assertEqual(len(rows), 5)

    def test_cli_generate(self):
        rc = run(["--config", str(self.dir / "config.toml"), "codes", "generate", "4"])
        self.assertEqual(rc, 0)
        self.assertEqual(len(read_table(self.cfg.path("codes"))), 9)  # 5 from the sample + 4

    def test_cli_rejects_zero(self):
        self.assertEqual(run(["--config", str(self.dir / "config.toml"), "codes", "generate", "0"]), 1)


class TestEmails(Base):
    def setUp(self):
        super().setUp()
        self.process()

    def draft(self, kind, eid):
        return (self.cfg.path("emails") / kind / f"{eid}.txt").read_text(encoding="utf-8")

    def test_received_for_everyone_and_nothing_left_unfilled(self):
        r = emails.make_drafts(self.cfg, "received")
        self.assertEqual(len(r["made"]), self.n)
        text = self.draft("received", "ZH-001")
        self.assertIn("To: ada@example.com", text)
        self.assertIn("Hi Ada,", text)
        self.assertIn("pharosprize@gmail.com", text)
        self.assertNotIn("[contact email]", text)
        self.assertIn("ZH-001", text)
        combined = (self.cfg.path("emails") / "received_ALL.txt").read_text(encoding="utf-8")
        self.assertIn(f"[ZH-{self.n:03d}]", combined)
        for eid in ("ZH-001", "ZH-013"):
            body = self.draft("received", eid)
            self.assertNotRegex(body.replace("[contact email]", ""), r"[{}]")

    def test_problem_drafts_use_the_entrants_flags(self):
        r = emails.make_drafts(self.cfg, "problem")
        self.assertNotIn("ZH-003", r["made"])  # exactly 800 words, nothing wrong
        self.assertNotIn("ZH-012", r["made"])
        self.assertIn("ZH-002", r["made"])
        self.assertIn("850 words", self.draft("problem", "ZH-002"))
        self.assertIn("scan or photo", self.draft("problem", "ZH-005"))
        self.assertIn("already been used", self.draft("problem", "ZH-014"))
        self.assertIn("could not recognise", self.draft("problem", "ZH-015"))
        self.assertIn("did not enter a fee support code", self.draft("problem", "ZH-016"))
        for eid in ("ZH-024", "ZH-025", "ZH-026", "ZH-028"):  # link, other link, text, .docx
            self.assertIn("isn't a PDF", self.draft("problem", eid))
            self.assertIn("attached as a PDF", self.draft("problem", eid))
        self.assertIn("could not find an essay file", self.draft("problem", "ZH-007"))
        self.assertIn("which entry fee option", self.draft("problem", "ZH-022"))
        multi = self.draft("problem", "ZH-021")
        for phrase in ("900 words", "could not read the age", "proof of your payment"):
            self.assertIn(phrase, multi)
        # nothing about other entrants leaks into an email
        self.assertNotRegex(self.draft("problem", "ZH-001"), r"ZH-008")
        self.assertNotRegex(self.draft("problem", "ZH-019"), r"ZH-020")

    def test_non_latin_name_email_says_entry_is_still_in_only_when_it_is_the_only_flag(self):
        asked = ("We print names on certificates in English letters, so could you send your name as you'd like "
                 "it to appear in English letters?")
        still_in = "Your entry is still in."
        emails.make_drafts(self.cfg, "problem")
        only = self.draft("problem", "ZH-023")        # non-Latin name and nothing else
        self.assertIn(asked + " " + still_in, only)
        others = self.draft("problem", "ZH-029")      # non-Latin name AND unreadable age
        self.assertIn(asked, others)
        self.assertNotIn(still_in, others)
        self.assertIn("age on your entry form", others)
        # fixing the other problem later flips the wording back
        rows = self.csv_rows()
        rows[28]["Age on 17 January 2027"] = "17"
        self.write_csv_rows(rows)
        self.process()
        emails.make_drafts(self.cfg, "problem", only={"ZH-029"})
        self.assertIn(still_in, self.draft("problem", "ZH-029"))

    def test_confirmed_needs_status_and_payment_verified(self):
        self.edit_tracker(clean={"Payment verified (Y/N)": "Y", "Status": "Confirmed"},
                          over={"Status": "Confirmed"})  # forgot to mark payment verified
        r = emails.make_drafts(self.cfg, "confirmed")
        self.assertEqual(r["made"], ["ZH-001"])
        self.assertTrue(any("ZH-002" in n for n in r["notes"]))
        self.assertIn("confirmed", self.draft("confirmed", "ZH-001"))

    def test_no_problem_email_once_decided_and_rejected_get_no_received(self):
        self.edit_tracker(over={"Status": "Confirmed", "Payment verified (Y/N)": "Y"}, old={"Status": "Rejected"})
        self.assertNotIn("ZH-002", emails.make_drafts(self.cfg, "problem")["made"])
        self.assertNotIn("ZH-011", emails.make_drafts(self.cfg, "problem")["made"])
        self.assertNotIn("ZH-011", emails.make_drafts(self.cfg, "received")["made"])

    def test_second_run_only_drafts_new_things(self):
        self.assertEqual(len(emails.make_drafts(self.cfg, "received")["made"]), self.n)
        again = emails.make_drafts(self.cfg, "received")
        self.assertEqual(again["made"], [])
        self.assertEqual(again["skipped_already_drafted"], self.n)
        self.assertEqual(len(emails.make_drafts(self.cfg, "received", redo_all=True)["made"]), self.n)
        self.assertEqual(emails.make_drafts(self.cfg, "received", only={"ZH-004"})["made"], ["ZH-004"])

    def test_problem_draft_redone_when_flags_change(self):
        emails.make_drafts(self.cfg, "problem")
        self.assertEqual(emails.make_drafts(self.cfg, "problem")["made"], [])
        # Jon fixes nothing, but a second, different problem appears: a new transfer-reference duplicate
        src = self.dir / "tally-export.csv"
        rows = self.csv_rows()
        rows[10]["Transfer reference"] = rows[0]["Transfer reference"]  # ZH-011 now duplicates ZH-001
        self.write_csv_rows(rows)
        self.process()
        made = emails.make_drafts(self.cfg, "problem")["made"]
        self.assertIn("ZH-011", made)
        self.assertIn("ZH-001", made)
        self.assertNotIn("ZH-002", made)

    def test_cli_sends_nothing_and_reports(self):
        self.assertEqual(run(["--config", str(self.dir / "config.toml"), "emails"]), 0)
        self.assertTrue((self.cfg.path("emails") / "problem_ALL.txt").exists())

    def test_templates_obey_content_red_lines(self):
        for f in (ROOT / "templates").iterdir():
            low = f.read_text(encoding="utf-8").lower()
            for word in FORBIDDEN:
                self.assertNotIn(word, low, f"{f.name} contains '{word}'")
        self.assertIn("run by student volunteers", (ROOT / "templates" / "received.txt").read_text())


class TestContactEmailMatchesSite(unittest.TestCase):
    def test_contact_email_matches_site_ts(self):
        """The address is typed in two places; fail if they ever drift apart. site.ts is only read."""
        site_ts = ROOT.parent.parent / "src" / "site.ts"
        if not site_ts.is_file():
            self.skipTest("src/site.ts not found (toolkit used outside the website repo)")
        m = re.search(r"""CONTACT_EMAIL\s*=\s*['"]([^'"]+)['"]""", site_ts.read_text(encoding="utf-8"))
        self.assertIsNotNone(m, "CONTACT_EMAIL not found in src/site.ts. Was it renamed? Update this test.")
        toolkit = tomllib.loads((ROOT / "config.toml").read_text(encoding="utf-8"))["email"]["contact_email"]
        self.assertEqual(toolkit, m.group(1),
                         "tools/entry-admin/config.toml [email] contact_email must match CONTACT_EMAIL in src/site.ts")


class TestSampleDataIsFake(unittest.TestCase):
    """The committed sample files must never hold real people's details."""

    def test_every_sample_person_is_fake(self):
        sample = ROOT / "sample"
        with (sample / "tally-export.csv").open(encoding="utf-8-sig", newline="") as f:
            rows = list(csv.DictReader(f))
        self.assertGreater(len(rows), 20)
        for r in rows:
            self.assertTrue(r["Email"].endswith("@example.com"), r["Email"])
            self.assertEqual(r["School"], "Fake Example School")
            self.assertTrue(r["Submission ID"].startswith("sub"))
        for p in (sample / "codes.xlsx",):
            for r in read_table(p):
                self.assertEqual((r["Issued to"], r["Email"]), ("", ""))

    def test_real_data_locations_are_git_ignored(self):
        import subprocess
        probes = ["work/submissions.csv", "work/tracker.xlsx", "work/judging/key.xlsx", "data/real/export.csv",
                  "real/export.csv", "downloads/essay.pdf", "export.csv", "tracker.xlsx", "judging/key.xlsx",
                  "emails/problem_ALL.txt", "sample/output/tracker.xlsx"]
        for rel in probes:
            r = subprocess.run(["git", "check-ignore", "-q", rel], cwd=ROOT, capture_output=True)
            if r.returncode == 128:
                self.skipTest("not inside a git checkout")
            self.assertEqual(r.returncode, 0, f"{rel} is not git-ignored")


class TestPack(Base):
    def setUp(self):
        super().setUp()
        self.process()

    def test_only_confirmed_and_verified_are_packed_with_anonymous_names(self):
        self.edit_tracker(clean={"Payment verified (Y/N)": "Y", "Status": "Confirmed"},
                          limit800={"Payment verified (Y/N)": "Y", "Status": "Confirmed"},
                          half_ok={"Payment verified (Y/N)": "Y", "Status": "Confirmed"},
                          over={"Status": "Confirmed"},                       # not verified: skipped
                          old={"Payment verified (Y/N)": "Y", "Status": "Rejected"})
        r = pack.pack(self.cfg)
        self.assertEqual(r["packed"], ["ZH-001", "ZH-003", "ZH-013"])
        out = self.cfg.path("judging")
        self.assertEqual(sorted(p.name for p in out.iterdir()), ["ZH-001.pdf", "ZH-003.pdf", "ZH-013.pdf", "key.xlsx"])
        self.assertTrue(any("ZH-002" in w for w in r["warnings"]))
        key = read_table(out / "key.xlsx")
        self.assertEqual([k["Entry ID"] for k in key], ["ZH-001", "ZH-003", "ZH-013"])
        self.assertEqual(key[0]["Full name"], "Ada Example")
        self.assertEqual(key[0]["Email"], "ada@example.com")

    def test_pdf_metadata_author_is_removed_and_text_kept(self):
        self.edit_tracker(clean={"Payment verified (Y/N)": "Y", "Status": "Confirmed"})
        src_author = PdfReader(str(self.dir / "downloads" / "essay-clean.pdf")).metadata.author
        self.assertEqual(src_author, "Ada Example")
        pack.pack(self.cfg)
        out = PdfReader(str(self.cfg.path("judging") / "ZH-001.pdf"))
        self.assertFalse(out.metadata and out.metadata.author)
        self.assertEqual(len(out.pages[0].extract_text().split()) > 10, True)

    def test_warns_when_essay_text_names_the_author(self):
        self.edit_tracker(half_ok={"Payment verified (Y/N)": "Y", "Status": "Confirmed"},
                          clean={"Payment verified (Y/N)": "Y", "Status": "Confirmed"})
        r = pack.pack(self.cfg)
        self.assertTrue(any("ZH-013.pdf" in w and "name" in w for w in r["warnings"]))
        self.assertFalse(any("ZH-001.pdf" in w for w in r["warnings"]))

    def test_entry_no_longer_confirmed_is_removed_from_judging(self):
        self.edit_tracker(clean={"Payment verified (Y/N)": "Y", "Status": "Confirmed"})
        pack.pack(self.cfg)
        self.assertTrue((self.cfg.path("judging") / "ZH-001.pdf").exists())
        self.edit_tracker(clean={"Status": "Rejected"})
        r = pack.pack(self.cfg)
        self.assertEqual(r["removed"], ["ZH-001.pdf"])
        self.assertFalse((self.cfg.path("judging") / "ZH-001.pdf").exists())

    def test_cli_pack(self):
        self.edit_tracker(clean={"Payment verified (Y/N)": "Y", "Status": "Confirmed"})
        self.assertEqual(run(["--config", str(self.dir / "config.toml"), "pack"]), 0)


if __name__ == "__main__":
    unittest.main()
