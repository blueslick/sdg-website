"""`process`: Tally CSV -> tracker.xlsx, with flags. Re-running keeps the owner's typed columns."""
from __future__ import annotations

import csv
import re
import unicodedata
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from . import codes as codes_mod
from .config import Config, ConfigError
from .pdfs import FileIndex, analyse
from .xlsxutil import OWNER_FILL, autosize, read_table, save_workbook, style_header

PAYMENT_COL = "Payment verified (Y/N)"
NOTES_COL = "Notes"
STATUS_COL = "Status"
OWNER_COLS = (PAYMENT_COL, NOTES_COL, STATUS_COL)
STATUSES = ("Pending", "Confirmed", "Rejected")

RED = PatternFill("solid", fgColor="F8CBAD")
AMBER = PatternFill("solid", fgColor="FFE699")

# key, column header, colour. Red = fix before confirming. Amber = look at it.
FLAGS = [
    ("over_word_limit", "Over word limit", RED),
    ("pdf_problem", "PDF problem", RED),
    ("duplicate_entrant", "Duplicate entrant", AMBER),
    ("name_not_latin", "Non-Latin name", AMBER),
    ("age_over_limit", "Age over limit", RED),
    ("fee_code", "Fee code problem", RED),
    ("no_proof", "No proof of payment", RED),
    ("duplicate_paypal", "Duplicate PayPal ID", RED),
]
FLAG_HEADERS = [h for _, h, _ in FLAGS]

DATA_COLS = ["Entry ID", "Submission ID", "Submitted at", "Full name", "Email", "Country", "School", "Age",
             "Prompt", "Fee", "Fee support code", "PayPal transaction ID", "Email paid with", "Declaration",
             "Essay file", "Word count"]
HEADERS = DATA_COLS + ["Flags"] + FLAG_HEADERS + list(OWNER_COLS)

# Columns the CSV must have. These two are allowed to be absent.
OPTIONAL_KEYS = {"respondent_id", "declaration"}


# ---------------------------------------------------------------- reading the CSV

def _norm_header(h: str) -> str:
    return " ".join(str(h).split()).casefold()


def read_csv(cfg: Config, path: Path) -> list[dict[str, str]]:
    """Read the Tally CSV into dicts keyed by the toolkit's own names (via [columns] in config)."""
    if not path.is_file():
        raise ConfigError(f"CSV not found: {path}\nPut the Tally export there, or pass --csv <file>.")
    raw = path.read_bytes()
    try:
        text = raw.decode("utf-8-sig")
    except UnicodeDecodeError:
        text = raw.decode("cp1252")  # a CSV re-saved by Excel on Windows
    lines = text.splitlines()
    if not lines:
        raise ConfigError(f"{path.name} is empty.")
    delim = max([",", ";", "\t"], key=lines[0].count)
    reader = csv.DictReader(text.splitlines(), delimiter=delim)
    actual = {_norm_header(h): h for h in (reader.fieldnames or [])}
    mapping: dict[str, str | None] = {}
    missing = []
    for key, title in cfg["columns"].items():
        found = actual.get(_norm_header(title))
        mapping[key] = found
        if found is None and key not in OPTIONAL_KEYS:
            missing.append((key, title))
    if missing:
        lines_ = "\n".join(f'  [columns] {k} = "{t}"' for k, t in missing)
        raise ConfigError(
            "These columns from config.toml are not in the CSV:\n" + lines_ +
            "\n\nColumns the CSV actually has:\n  " + "\n  ".join(reader.fieldnames or []) +
            "\n\nIf you renamed a question in Tally, change the text after = in config.toml.")
    rows = []
    for r in reader:
        row = {k: (r.get(col) or "").strip() if col else "" for k, col in mapping.items()}
        if row["submission_id"]:
            rows.append(row)
    seen: set[str] = set()
    unique = []
    for r in rows:  # same Submission ID twice in one export: keep the first
        if r["submission_id"] not in seen:
            seen.add(r["submission_id"])
            unique.append(r)
    return unique


def parse_time(s: str) -> datetime:
    s = s.strip().replace("Z", "+00:00")
    for parse in (datetime.fromisoformat,
                  lambda x: datetime.strptime(x, "%d/%m/%Y %H:%M:%S"),
                  lambda x: datetime.strptime(x, "%d/%m/%Y %H:%M"),
                  lambda x: datetime.strptime(x, "%m/%d/%Y %H:%M:%S")):
        try:
            return parse(s).replace(tzinfo=None)
        except ValueError:
            continue
    return datetime.min


def parse_age(s: str) -> int | None:
    m = re.search(r"\d+", s or "")
    return int(m.group()) if m else None


def norm_name(s: str) -> str:
    return " ".join(re.sub(r"[^\w\s]", "", s.casefold()).split())


def non_latin_chars(name: str) -> str:
    """Letters in the name that are not Latin script (e.g. Chinese, Arabic, Cyrillic). Accented Latin letters are fine."""
    bad = [ch for ch in name if ch.isalpha() and "LATIN" not in unicodedata.name(ch, "")]
    return "".join(bad)


def norm_fee(s: str) -> str:
    return " ".join(s.casefold().split())


def norm_txn(s: str) -> str:
    return "".join(ch for ch in s.upper() if ch.isalnum())


# ---------------------------------------------------------------- flags

def _groups(entries: list[dict], keyfn) -> dict[str, list[dict]]:
    g: dict[str, list[dict]] = {}
    for e in entries:
        k = keyfn(e)
        if k:
            g.setdefault(k, []).append(e)
    return {k: v for k, v in g.items() if len(v) > 1}


def compute_flags(entries: list[dict], cfg: Config, code_table: dict[str, dict]) -> dict[str, str]:
    """Fills e['flags'] (key -> text) for every entry. Returns {normalised code: first user's entry ID}."""
    ec = cfg["entry"]
    full_fee, half_fee = norm_fee(cfg["fees"]["full"]), norm_fee(cfg["fees"]["half"])

    for e in entries:
        e["flags"] = {}

    # essay / word count
    claimed: dict[Path, list[dict]] = {}
    for e in entries:
        f = e["flags"]
        res = e.get("pdf")
        if not e["essay_cell"]:
            f["pdf_problem"] = "Missing: no essay uploaded"
        elif e["essay_path"] is None:
            f["pdf_problem"] = f"Missing: '{e['essay_wanted']}' not found in the downloads folder"
        else:
            claimed.setdefault(e["essay_path"], []).append(e)
            if res.status == "unreadable":
                f["pdf_problem"] = f"Unreadable: {res.detail}"
            elif res.status == "scanned":
                f["pdf_problem"] = f"Scanned: {res.detail}"
            elif res.words > ec["word_limit"]:
                f["over_word_limit"] = f"{res.words} words (limit {ec['word_limit']})"
    for path, group in claimed.items():
        if len(group) > 1:  # two entries matched the same downloaded file: one of them is wrong
            for e in group:
                others = ", ".join(o["entry_id"] for o in group if o is not e)
                e["flags"]["pdf_problem"] = (f"Shared: same downloaded file as {others}; "
                                             "check the downloads (see README)")

    # duplicates
    by_email = _groups(entries, lambda e: e["email"].strip().casefold())
    by_name = _groups(entries, lambda e: norm_name(e["full_name"]))
    for kind, groups in (("Email", by_email), ("Name", by_name)):
        for group in groups.values():
            for e in group:
                others = ", ".join(o["entry_id"] for o in group if o is not e)
                prev = e["flags"].get("duplicate_entrant")
                text = f"{kind}: same as {others}"
                e["flags"]["duplicate_entrant"] = f"{prev}; {text}" if prev else text
    for group in _groups(entries, lambda e: norm_txn(e["paypal_txn"])).values():
        for e in group:
            others = ", ".join(o["entry_id"] for o in group if o is not e)
            e["flags"]["duplicate_paypal"] = f"Same transaction ID as {others}"

    # name must be in English (Latin) letters
    for e in entries:
        bad = non_latin_chars(e["full_name"])
        if bad:
            e["flags"]["name_not_latin"] = f"Non-Latin characters: {bad}"

    # age
    for e in entries:
        age = parse_age(e["age"])
        if age is None:
            e["flags"]["age_over_limit"] = f"Unclear: age is '{e['age']}'" if e["age"] else "Unclear: no age given"
        elif age > ec["max_age"]:
            e["flags"]["age_over_limit"] = f"Over: {age} (limit {ec['max_age']})"

    # proof of payment
    for e in entries:
        if not e["proof_cell"]:
            extra = " (a transaction ID was given)" if e["paypal_txn"] else ""
            e["flags"]["no_proof"] = f"Missing: no file uploaded{extra}"

    # fee answer, and for half fee the support code. The first entry (by time) to use a valid code owns it,
    # unless codes.xlsx already says someone used it.
    first_user: dict[str, str] = {}
    for e in sorted(entries, key=lambda e: (e["time"], e["order"])):
        fee = norm_fee(e["fee"])
        if fee == full_fee:
            continue
        if fee != half_fee:  # not one of the two options on the form: don't guess
            shown = e["fee"] or "(blank)"
            e["flags"]["fee_code"] = f"Unclear: fee answer '{shown}' is not one of the two options in config.toml"
            continue
        code = codes_mod.normalise(e["fee_code"])
        if not code:
            e["flags"]["fee_code"] = "Missing: half fee chosen but no code entered"
        elif code not in code_table:
            e["flags"]["fee_code"] = f"Invalid: '{e['fee_code']}' is not on the codes list"
        else:
            recorded = str(code_table[code].get("Used by submission") or "").strip()
            owner = recorded or first_user.get(code, "")
            if not owner:
                first_user[code] = e["entry_id"]
            elif owner not in (e["entry_id"], e["submission_id"]):
                e["flags"]["fee_code"] = f"Used: already used by {owner}"
    return first_user


# ---------------------------------------------------------------- building the tracker

def _assign_ids(csv_rows: list[dict], old: dict[str, dict], prefix: str, digits: int) -> None:
    pat = re.compile(rf"^{re.escape(prefix)}-(\d+)$")
    nums = [int(m.group(1)) for r in old.values() if (m := pat.match(str(r.get("Entry ID", ""))))]
    nxt = max(nums, default=0) + 1
    for r in csv_rows:
        known = old.get(r["submission_id"])
        if known and known.get("Entry ID"):
            r["entry_id"] = str(known["Entry ID"])
        else:
            r["entry_id"] = f"{prefix}-{nxt:0{digits}d}"
            r["is_new"] = True
            nxt += 1


def process(cfg: Config, csv_path: Path | None = None, files_dir: Path | None = None) -> dict:
    ec = cfg["entry"]
    csv_path = csv_path or cfg.path("csv")
    files_dir = files_dir or cfg.path("files")
    tracker_path = cfg.path("tracker")
    codes_path = cfg.path("codes")
    notes: list[str] = []

    rows = read_csv(cfg, csv_path)
    if not files_dir.is_dir():
        notes.append(f"Downloads folder not found ({files_dir}), so every essay will be flagged as missing.")
    if not codes_path.exists():
        notes.append(f"No codes file at {codes_path}, so every half-fee code will look invalid. "
                     "Run `codes generate N` first.")

    old_rows = read_table(tracker_path)
    old = {str(r.get("Submission ID")): r for r in old_rows if r.get("Submission ID")}
    for i, r in enumerate(rows):
        r["time"] = parse_time(r["submitted_at"])
        r["order"] = i
    rows.sort(key=lambda r: (r["time"], r["order"]))
    _assign_ids(rows, old, ec["id_prefix"], ec["id_digits"])

    index = FileIndex(files_dir)
    entries = []
    for r in rows:
        path, wanted = index.find(r["essay"], r["submission_id"]) if r["essay"] else (None, "")
        pdf = analyse(path, ec["min_words_for_text"]) if path else None
        entries.append({
            "entry_id": r["entry_id"], "submission_id": r["submission_id"], "time": r["time"], "order": r["order"],
            "submitted_at": r["submitted_at"], "full_name": r["full_name"], "email": r["email"],
            "country": r["country"], "school": r["school"], "age": r["age"], "prompt": r["prompt"],
            "fee": r["fee"], "fee_code": r["fee_code"], "paypal_txn": r["paypal_txn"],
            "paypal_email": r["paypal_email"], "declaration": r["declaration"],
            "essay_cell": r["essay"], "essay_path": path, "essay_wanted": wanted, "pdf": pdf,
            "proof_cell": r["proof"], "is_new": r.get("is_new", False),
        })

    code_table = codes_mod.load_codes(codes_path)
    first_user = compute_flags(entries, cfg, code_table)
    if first_user and codes_path.exists():
        codes_mod.record_usage(codes_path, first_user)

    # Rows already in the tracker but no longer in the CSV are kept exactly as they were.
    in_csv = {e["submission_id"] for e in entries}
    orphans = [r for sid, r in old.items() if sid not in in_csv]
    for r in orphans:
        notes.append(f"{r.get('Entry ID')} ({r.get('Full name')}) is in the tracker but not in this CSV; kept as it was.")

    known_headers = set(HEADERS)
    extra_headers = [h for r in old_rows for h in r if h not in known_headers]
    extra_headers = list(dict.fromkeys(extra_headers))

    table: list[list] = []
    flag_cells: list[dict[str, str]] = []
    for e in entries:
        o = old.get(e["submission_id"], {})
        res = e["pdf"]
        age = parse_age(e["age"])
        flags = e["flags"]
        values = {
            "Entry ID": e["entry_id"], "Submission ID": e["submission_id"], "Submitted at": e["submitted_at"],
            "Full name": e["full_name"], "Email": e["email"], "Country": e["country"], "School": e["school"],
            "Age": age if age is not None and e["age"].strip().isdigit() else e["age"],
            "Prompt": e["prompt"], "Fee": e["fee"], "Fee support code": e["fee_code"],
            "PayPal transaction ID": e["paypal_txn"], "Email paid with": e["paypal_email"],
            "Declaration": e["declaration"],
            "Essay file": e["essay_path"].relative_to(files_dir).as_posix() if e["essay_path"] else "",
            "Word count": res.words if res and res.status != "unreadable" else "",
            "Flags": len(flags),
        }
        for key, header, _ in FLAGS:
            values[header] = flags.get(key, "")
        for h in list(OWNER_COLS) + extra_headers:
            values[h] = o.get(h, "")
        table.append([values.get(h, "") for h in HEADERS + extra_headers])
        flag_cells.append(flags)
    for r in orphans:
        table.append([r.get(h, "") for h in HEADERS + extra_headers])
        flag_cells.append({k: r[h] for k, h, _ in FLAGS if r.get(h)})

    order = sorted(range(len(table)), key=lambda i: table[i][0])
    table = [table[i] for i in order]
    flag_cells = [flag_cells[i] for i in order]
    _write(tracker_path, table, flag_cells, extra_headers)

    return {"total": len(table), "new": [e["entry_id"] for e in entries if e["is_new"]],
            "flagged": sum(1 for f in flag_cells if f), "notes": notes,
            "flag_counts": {h: sum(1 for f in flag_cells if k in f) for k, h, _ in FLAGS},
            "path": tracker_path}


def _write(path: Path, table: list[list], flag_cells: list[dict], extra: list[str]) -> None:
    headers = HEADERS + extra
    wb = Workbook()
    ws = wb.active
    ws.title = "Entries"
    style_header(ws, headers, owner_cols=OWNER_COLS)
    for row in table:
        ws.append(row)
    for r_idx, flags in enumerate(flag_cells, start=2):
        for key, header, fill in FLAGS:
            c = ws.cell(row=r_idx, column=headers.index(header) + 1)
            if key in flags or c.value:
                c.fill = fill
                c.alignment = Alignment(wrap_text=True, vertical="top")
    autosize(ws, headers)
    for h in FLAG_HEADERS:
        ws.column_dimensions[get_column_letter(headers.index(h) + 1)].width = 28
    last = max(len(table) + 1, 2) + 500  # drop-downs and colours also cover rows added later by hand
    pay = get_column_letter(headers.index(PAYMENT_COL) + 1)
    st = get_column_letter(headers.index(STATUS_COL) + 1)
    dv_pay = DataValidation(type="list", formula1='"Y,N"', allow_blank=True)
    dv_st = DataValidation(type="list", formula1='"' + ",".join(STATUSES) + '"', allow_blank=True)
    ws.add_data_validation(dv_pay)
    ws.add_data_validation(dv_st)
    dv_pay.add(f"{pay}2:{pay}{last}")
    dv_st.add(f"{st}2:{st}{last}")
    green = PatternFill("solid", start_color="C6E0B4", end_color="C6E0B4")
    grey = PatternFill("solid", start_color="D9D9D9", end_color="D9D9D9")
    ws.conditional_formatting.add(f"{st}2:{st}{last}", CellIsRule(operator="equal", formula=['"Confirmed"'], fill=green))
    ws.conditional_formatting.add(f"{st}2:{st}{last}", CellIsRule(operator="equal", formula=['"Rejected"'], fill=grey,
                                                                   font=Font(color="595959")))
    ws.conditional_formatting.add(f"{pay}2:{pay}{last}", CellIsRule(operator="equal", formula=['"Y"'], fill=green))
    legend = wb.create_sheet("Key")
    legend.append(["Colour", "Meaning"])
    legend.append(["Red cell", "Fix or ask the entrant before confirming."])
    legend.append(["Amber cell", "Look at it. Might be fine (e.g. someone who submitted twice by accident)."])
    legend.append(["Yellow header", "Your columns. Type here. Everything else is rewritten from the CSV each time you run `process`."])
    legend["A2"].fill, legend["A3"].fill, legend["A4"].fill = RED, AMBER, OWNER_FILL
    legend.column_dimensions["A"].width = 16
    legend.column_dimensions["B"].width = 110
    save_workbook(wb, path, keep_backup=True)


# ---------------------------------------------------------------- reading it back (emails, pack)

def load_tracker(cfg: Config) -> list[dict]:
    path = cfg.path("tracker")
    if not path.exists():
        raise ConfigError(f"No tracker at {path}. Run `process` first.")
    return read_table(path)


def status_of(row: dict) -> str:
    s = str(row.get(STATUS_COL) or "").strip().title()
    return s if s in STATUSES else "Pending"


def raised_flags(row: dict) -> list[tuple[str, str]]:
    """[(flag key, text)] for the flags set on a tracker row."""
    return [(k, str(row[h]).strip()) for k, h, _ in FLAGS if str(row.get(h) or "").strip()]


def is_ready(row: dict) -> tuple[bool, str]:
    """Confirmed AND payment verified = ready for the confirmed email / judging pack."""
    if status_of(row) != "Confirmed":
        return False, ""
    if str(row.get(PAYMENT_COL) or "").strip().upper() != "Y":
        return False, f"{row.get('Entry ID')} is Confirmed but '{PAYMENT_COL}' is not Y, so it was skipped."
    return True, ""
