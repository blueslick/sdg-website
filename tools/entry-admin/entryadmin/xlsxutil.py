"""Small helpers for reading and writing Excel files safely."""
from __future__ import annotations

import os
import shutil
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .config import ConfigError

HEADER_FILL = PatternFill("solid", fgColor="DDDDDD")
OWNER_FILL = PatternFill("solid", fgColor="FFF2CC")
HEADER_FONT = Font(bold=True)


def read_table(path: Path) -> list[dict[str, object]]:
    """Read a sheet whose first row is headers into a list of dicts. Empty cells -> ''."""
    if not path.exists():
        return []
    try:
        wb = load_workbook(path, data_only=True)
    except PermissionError as e:
        raise ConfigError(f"Can't open {path.name}. Is it open in Excel? Close it and try again.") from e
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return []
    headers = [str(h).strip() if h is not None else "" for h in rows[0]]
    out = []
    for r in rows[1:]:
        if all(v is None or v == "" for v in r):
            continue
        out.append({h: ("" if v is None else v) for h, v in zip(headers, r) if h})
    return out


def style_header(ws, headers: list[str], owner_cols: tuple[str, ...] = ()) -> None:
    for i, h in enumerate(headers, start=1):
        c = ws.cell(row=1, column=i, value=h)
        c.font = HEADER_FONT
        c.fill = OWNER_FILL if h in owner_cols else HEADER_FILL
        c.alignment = Alignment(wrap_text=True, vertical="center")
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}1"


def autosize(ws, headers: list[str], max_width: int = 45) -> None:
    for i, h in enumerate(headers, start=1):
        longest = max([len(str(h))] + [len(str(c.value)) for c in ws[get_column_letter(i)][1:] if c.value is not None])
        ws.column_dimensions[get_column_letter(i)].width = min(max(longest + 2, 10), max_width)


def save_workbook(wb: Workbook, path: Path, keep_backup: bool = False) -> None:
    """Write via a temp file so a failure can't leave a half-written file.

    keep_backup copies the old file to '<name>.previous.xlsx' first, a safety net for files
    the owner types in.
    """
    path.parent.mkdir(parents=True, exist_ok=True)
    tmp = path.with_name(path.name + ".tmp")
    wb.save(tmp)
    try:
        if keep_backup and path.exists():
            shutil.copy2(path, path.with_name(path.stem + ".previous.xlsx"))
        os.replace(tmp, path)
    except PermissionError as e:
        tmp.unlink(missing_ok=True)
        raise ConfigError(f"Can't write {path.name}. Is it open in Excel? Close it and try again.") from e
