"""Make the scoring sheet: one Excel file, one row per essay, for the one judge (the owner).

The sheet only ever holds entry IDs, never names: this command does not read key.xlsx.

Columns:  Entry ID | Triage | one column per criterion | Total | Comments | Second reader
  Triage         Clearly no / Maybe / Strong, for every essay.
  Criteria       whole marks, up to the marks for that criterion. Only for Maybe and Strong.
  Total          adds itself up once every criterion has a mark.
  Second reader  optional, only for Clearly-no essays: Agree / Read again (the spot check).
"""
from __future__ import annotations

import json
from datetime import datetime
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill, Protection
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from common import JudgingError, Tokens
from rubric import Rubric

META_SHEET = "_meta"
SCORES_SHEET = "Scores"
TRIAGE_NO, TRIAGE_MAYBE, TRIAGE_STRONG = "Clearly no", "Maybe", "Strong"
TRIAGE = (TRIAGE_NO, TRIAGE_MAYBE, TRIAGE_STRONG)
SECOND_AGREE, SECOND_AGAIN = "Agree", "Read again"
SECOND = (SECOND_AGREE, SECOND_AGAIN)


def layout(rubric: Rubric) -> dict[str, int]:
    """1-based column numbers, shared by the sheet maker and the results reader."""
    n = len(rubric.criteria)
    return {"id": 1, "triage": 2, "first_score": 3, "last_score": 2 + n,
            "total": 3 + n, "comments": 4 + n, "second": 5 + n}


def _list_rule(options: tuple[str, ...], title: str, prompt: str) -> DataValidation:
    return DataValidation(type="list", formula1='"' + ",".join(options) + '"', allow_blank=True,
                          showErrorMessage=True, errorStyle="stop", errorTitle=f"{title} not allowed",
                          error="Pick one from the list: " + ", ".join(options) + ".",
                          showInputMessage=True, promptTitle=title, prompt=prompt)


def build_sheet(path: Path, entry_ids: list[str], rubric: Rubric, tokens: Tokens | None) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = SCORES_SHEET
    col = layout(rubric)
    n = len(rubric.criteria)

    def colour(name: str, fallback: str) -> str:
        return (tokens.color_or(name, fallback) if tokens else fallback).lstrip("#")

    head_fill = PatternFill("solid", fgColor=colour("cream", "FBEEE8"))
    head_font = Font(bold=True, color=colour("plum", "3A2226"))
    headers = ["Entry ID", "Triage", *[c.header() for c in rubric.criteria],
               f"Total (out of {rubric.total})", "Comments", "Second reader (Clearly no only, optional)"]
    for i, text in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=i, value=text)
        cell.font, cell.fill = head_font, head_fill
        cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="left" if i in (1, col["comments"]) else "center")
    ws.row_dimensions[1].height = 48

    first, last = 2, len(entry_ids) + 1
    unlocked = Protection(locked=False)
    f_col, l_col = get_column_letter(col["first_score"]), get_column_letter(col["last_score"])
    for i, entry_id in enumerate(entry_ids):
        row = first + i
        ws.cell(row=row, column=1, value=entry_id).font = Font(name="Consolas", bold=True)
        for c in (col["triage"], *range(col["first_score"], col["last_score"] + 1), col["comments"], col["second"]):
            cell = ws.cell(row=row, column=c)
            cell.protection = unlocked
            cell.alignment = Alignment(horizontal="left" if c == col["comments"] else "center",
                                       wrap_text=c == col["comments"], vertical="top" if c == col["comments"] else None)
        ws.cell(row=row, column=col["total"],
                value=f'=IF(COUNT({f_col}{row}:{l_col}{row})={n},SUM({f_col}{row}:{l_col}{row}),"")'
                ).alignment = Alignment(horizontal="center")

    triage = _list_rule(TRIAGE, "Triage", "Clearly no, Maybe or Strong")
    triage.add(f"{get_column_letter(col['triage'])}{first}:{get_column_letter(col['triage'])}{last}")
    second = _list_rule(SECOND, "Second reader", "Only for Clearly-no essays: Agree or Read again")
    second.add(f"{get_column_letter(col['second'])}{first}:{get_column_letter(col['second'])}{last}")
    ws.add_data_validation(triage)
    ws.add_data_validation(second)
    for i, criterion in enumerate(rubric.criteria):
        letter = get_column_letter(col["first_score"] + i)
        rule = DataValidation(type="whole", operator="between", formula1="0", formula2=str(criterion.marks),
                              allow_blank=True, showErrorMessage=True, errorStyle="stop",
                              errorTitle="Mark not allowed",
                              error=f"Enter a whole number from 0 to {criterion.marks}.",
                              showInputMessage=True, promptTitle=criterion.name,
                              prompt=f"Whole marks, 0 to {criterion.marks}")
        rule.add(f"{letter}{first}:{letter}{last}")
        ws.add_data_validation(rule)

    ws.column_dimensions["A"].width = 12
    ws.column_dimensions[get_column_letter(col["triage"])].width = 14
    for c in range(col["first_score"], col["last_score"] + 1):
        ws.column_dimensions[get_column_letter(c)].width = 18
    ws.column_dimensions[get_column_letter(col["total"])].width = 14
    ws.column_dimensions[get_column_letter(col["comments"])].width = 60
    ws.column_dimensions[get_column_letter(col["second"])].width = 22
    ws.freeze_panes = "C2"
    # Lock the Entry ID and Total columns so they cannot be changed by accident. No password:
    # this guards against slips, it is not security.
    ws.protection.sheet = True
    ws.protection.formatColumns = False
    ws.protection.formatRows = False

    guide = wb.create_sheet("How to score")
    lines = [
        ("How to score", True),
        ("1. Triage every essay: Clearly no, Maybe or Strong.", False),
        (f"2. For Maybe and Strong essays, give a whole mark for every criterion (up to its maximum). "
         f"The Total is out of {rubric.total}.", False),
        ("3. Optional spot check: a second reader can look at some Clearly-no essays and write Agree or "
         "Read again. Read again is listed in results.xlsx so you can look at it once more.", False),
        ("Scores are given by people only. The toolkit adds up your marks and sorts them; it never picks "
         "the shortlist or the winners.", False),
        ("The wording of each criterion is in the judging criteria on the Competition page.", False),
        ("", False),
        ("Criterion", True),
    ]
    if rubric.placeholder:
        lines.insert(1, ("NOTE: the rubric is still marked as a PLACEHOLDER in rubric.yaml.", True))
    for text, bold in lines:
        guide.append([text])
        guide.cell(row=guide.max_row, column=1).font = Font(bold=bold)
    guide.cell(row=guide.max_row, column=2, value="Marks").font = Font(bold=True)
    for c in rubric.criteria:
        guide.append([c.name, c.marks])
    guide.append(["Total", rubric.total])
    guide.column_dimensions["A"].width = 110

    meta = wb.create_sheet(META_SHEET)
    meta["A1"] = json.dumps({
        "version": 2, "entries": entry_ids,
        "criteria": [{"id": c.id, "marks": c.marks} for c in rubric.criteria],
        "created": datetime.now().isoformat(timespec="seconds"),
    })
    meta.sheet_state = "hidden"
    wb.active = 0
    wb.save(path)


def make_sheet(entry_ids: list[str], rubric: Rubric, path: Path, tokens: Tokens | None, force: bool = False) -> None:
    if path.exists() and not force:
        raise JudgingError(
            f"{path} already exists, and may hold your scores. Nothing was changed. "
            "Use --force only if you really want to start a new, empty sheet (keep a copy of the old one first)."
        )
    path.parent.mkdir(parents=True, exist_ok=True)
    build_sheet(path, sorted(entry_ids), rubric, tokens)
