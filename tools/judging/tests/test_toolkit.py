"""Tests for the judging toolkit. Run from anywhere:  python -m pytest tools/judging
They use only fake data (demo.py) and never touch tools/judging/data/ or src/data/certificates.json."""
from __future__ import annotations

import contextlib
import io
import json
import re
import shutil
import subprocess
from pathlib import Path

import openpyxl
import pytest

import demo
import judging
from certificates import EDITION, LEDGER_FILE, Recipient, issue_certificates, read_recipients, statement
from common import CATEGORIES, TOOL_DIR, JudgingError, Tokens, find_repo_root, read_key, site_values
from rubric import load_rubric

REPO = find_repo_root()
SITE = site_values(REPO)
PNPM_MISSING = not (REPO / "node_modules" / "@fontsource" / "gloock" / "files").is_dir()
needs_fonts = pytest.mark.skipif(PNPM_MISSING, reason="run `npm install` in the repo root first (the certificate fonts come from it)")


def run(*argv: str) -> tuple[int, str]:
    out, err = io.StringIO(), io.StringIO()
    with contextlib.redirect_stdout(out), contextlib.redirect_stderr(err):
        code = judging.main(list(argv))
    return code, out.getvalue() + err.getvalue()


def rows(path: Path, sheet: str) -> list[dict]:
    ws = openpyxl.load_workbook(path, data_only=True)[sheet]
    it = ws.iter_rows(values_only=True)
    header = [str(h) for h in next(it)]
    return [dict(zip(header, r)) for r in it if any(v is not None for v in r)]


# --------------------------------------------------------------------------- the fake competition

@pytest.fixture(scope="module")
def comp(tmp_path_factory):
    """Fake essays -> triaged scoring sheet -> results.xlsx (before any shortlist or winner is chosen)."""
    root = tmp_path_factory.mktemp("demo")
    demo.make_inputs(root)
    assert run("sheet", "--essays", str(root / "essays"), "--out", str(root / "scoring.xlsx"))[0] == 0
    demo.fill_in_scoring(root / "scoring.xlsx")
    code, out = run("results", "--scores", str(root / "scoring.xlsx"), "--key", str(root / "key.xlsx"),
                    "--essays", str(root / "essays"), "--out", str(root / "results.xlsx"))
    assert code == 0, out
    return root


@pytest.fixture(scope="module")
def fonts():
    if PNPM_MISSING:
        pytest.skip("needs `npm install`")
    from fontkit import setup_fonts
    setup_fonts(Tokens(REPO))


# --------------------------------------------------------------------------- the rubric

def test_rubric_yaml_matches_the_site_rubric():
    """If src/site.ts RUBRIC and rubric.yaml ever disagree (names or marks, or their order), this fails."""
    rubric = load_rubric(TOOL_DIR / "rubric.yaml")
    assert [(c.name, c.marks) for c in rubric.criteria] == SITE["rubric"]
    assert rubric.total == SITE["rubric_total"] == sum(m for _, m in SITE["rubric"]) == 100
    assert SITE["rubric"] == [("Response to the prompt", 30), ("Insight and understanding", 30),
                              ("Clarity and structure", 20), ("Voice and originality", 20)]
    assert rubric.placeholder is False


def test_rubric_marks_must_add_up(tmp_path):
    bad = tmp_path / "rubric.yaml"
    bad.write_text("total: 100\ncriteria:\n  - {id: a, name: A, marks: 60}\n  - {id: b, name: B, marks: 30}\n")
    with pytest.raises(JudgingError, match="add up to 90"):
        load_rubric(bad)
    bad.write_text("total: 100\ncriteria:\n  - {id: a, name: A, marks: 99.5}\n  - {id: b, name: B, marks: 0.5}\n")
    with pytest.raises(JudgingError, match="whole number"):
        load_rubric(bad)


def test_categories_and_edition_match_the_site():
    assert tuple(SITE["categories"]) == CATEGORIES == ("Junior", "Open")
    assert EDITION == "Inaugural Edition: Zero Hunger"
    assert EDITION in SITE["text"]                      # the site says the same words (OG_IMAGE_ALT)


# --------------------------------------------------------------------------- scoring sheet and results

def test_scoring_sheet_has_triage_marks_and_optional_second_reader(comp):
    ws = openpyxl.load_workbook(comp / "scoring.xlsx")["Scores"]
    header = [c.value for c in ws[1]]
    assert header[:2] == ["Entry ID", "Triage"]
    assert header[2:6] == ["Response to the prompt (out of 30)", "Insight and understanding (out of 30)",
                           "Clarity and structure (out of 20)", "Voice and originality (out of 20)"]
    assert header[6].startswith("Total") and header[7] == "Comments" and header[8].startswith("Second reader")
    kinds = {v.type: v.formula1 for v in ws.data_validations.dataValidation}
    assert kinds["list"] in ('"Clearly no,Maybe,Strong"', '"Agree,Read again"')
    limits = sorted(v.formula2 for v in ws.data_validations.dataValidation if v.type == "whole")
    assert limits == ["20", "20", "30", "30"]
    # No names anywhere in the sheet.
    names = {n for _, n, _ in demo.FAKE}
    text = " ".join(str(c.value) for sheet in openpyxl.load_workbook(comp / "scoring.xlsx") for r in sheet.iter_rows() for c in r)
    assert not any(n in text for n in names)


def test_sheet_will_not_overwrite_an_existing_one(comp):
    code, out = run("sheet", "--essays", str(comp / "essays"), "--out", str(comp / "scoring.xlsx"))
    assert code == 1 and "already exists" in out


def test_results_rank_junior_and_open_separately(comp):
    junior, open_ = rows(comp / "results.xlsx", "Junior"), rows(comp / "results.xlsx", "Open")
    assert {r["Entry ID"] for r in junior} <= {f"ZH-{i:03d}" for i in range(1, 9)}
    assert {r["Entry ID"] for r in open_} <= {f"ZH-{i:03d}" for i in range(9, 19)}
    assert [(r["Entry ID"], r["Rank"], r["Total (out of 100)"]) for r in junior[:3]] == [
        ("ZH-001", 1, 88), ("ZH-002", 2, 87), ("ZH-003", 3, 80)]
    assert [(r["Entry ID"], r["Rank"], r["Total (out of 100)"]) for r in open_[:3]] == [
        ("ZH-009", 1, 94), ("ZH-010", 2, 88), ("ZH-011", 3, 83)]       # Open's rank 1 is not mixed with Junior's
    assert junior[0]["Name"] == "Test Student 01" and junior[0]["Country"] == "Testland"
    assert [r["Total (out of 100)"] for r in junior] == sorted((r["Total (out of 100)"] for r in junior), reverse=True)


def test_ties_are_shown_not_broken(comp):
    junior = rows(comp / "results.xlsx", "Junior")
    tied = [r for r in junior if r["Total (out of 100)"] == 77]
    assert [r["Entry ID"] for r in tied] == ["ZH-004", "ZH-005"]
    assert [r["Rank"] for r in tied] == [4, 4]                        # shared rank
    assert all("Tie" in r["Flags"] for r in tied)
    assert [r["Rank"] for r in junior] == [1, 2, 3, 4, 4]
    assert any(f["Type"] == "Tie" and "ZH-004" in f["Detail"] and "ZH-005" in f["Detail"] for f in rows(comp / "results.xlsx", "Flags"))


def test_the_tool_never_picks_the_shortlist_or_winners(comp):
    for category in CATEGORIES:
        for r in rows(comp / "results.xlsx", category):
            assert r["Shortlist (yes/no)"] is None and r["Winner (yes/no)"] is None


def test_triage_rules(comp):
    flags = rows(comp / "results.xlsx", "Flags")
    kinds = {(f["Type"], f["Entry ID"]) for f in flags}
    assert ("Read again", "ZH-007") in kinds                          # second reader said Read again
    assert ("Not triaged", "ZH-018") in kinds
    assert ("Incomplete marks", "ZH-014") in kinds
    not_ranked = {r["Entry ID"]: r for r in rows(comp / "results.xlsx", "Not ranked")}
    assert not_ranked["ZH-006"]["Why it is not ranked"] == "Clearly no" and not_ranked["ZH-006"]["Second reader"] == "Agree"
    assert not_ranked["ZH-007"]["Second reader"] == "Read again"
    assert not_ranked["ZH-014"]["Why it is not ranked"] == "Marks missing"
    assert not_ranked["ZH-018"]["Why it is not ranked"] == "Not triaged yet"
    ranked = {r["Entry ID"] for c in CATEGORIES for r in rows(comp / "results.xlsx", c)}
    assert ranked.isdisjoint(not_ranked) and len(ranked) + len(not_ranked) == 18


def test_marks_on_a_clearly_no_essay_are_ignored_and_bad_marks_are_flagged(comp, tmp_path):
    sheet = tmp_path / "scoring.xlsx"
    shutil.copy(comp / "scoring.xlsx", sheet)
    wb = openpyxl.load_workbook(sheet)
    ws = wb["Scores"]
    for row in range(2, ws.max_row + 1):
        entry = ws.cell(row=row, column=1).value
        if entry == "ZH-008":                       # Clearly no, but marks typed in
            for c in range(3, 7):
                ws.cell(row=row, column=c, value=10)
        if entry == "ZH-001":                       # a pasted mark that bypassed the dropdown rules
            ws.cell(row=row, column=3, value=31)
    wb.save(sheet)
    code, out = run("results", "--scores", str(sheet), "--key", str(comp / "key.xlsx"), "--out", str(tmp_path / "r.xlsx"))
    assert code == 0
    kinds = {(f["Type"], f["Entry ID"]) for f in rows(tmp_path / "r.xlsx", "Flags")}
    assert ("Marks ignored", "ZH-008") in kinds and ("Invalid mark", "ZH-001") in kinds and ("Incomplete marks", "ZH-001") in kinds
    assert "ZH-008" not in {r["Entry ID"] for r in rows(tmp_path / "r.xlsx", "Junior")}


def test_key_needs_a_valid_category(tmp_path):
    wb = openpyxl.Workbook()
    wb.active.append(["Entry ID", "Full name", "Age"])
    wb.active.append(["ZH-001", "Test Student 01", 12])
    wb.save(tmp_path / "nocat.xlsx")
    with pytest.raises(JudgingError, match="Category"):
        read_key(tmp_path / "nocat.xlsx")
    wb = openpyxl.Workbook()
    wb.active.append(["Entry ID", "Full name", "Category"])
    wb.active.append(["ZH-001", "Test Student 01", "junior"])
    wb.active.append(["ZH-002", "Test Student 02", "Senior"])
    wb.save(tmp_path / "bad.xlsx")
    with pytest.raises(JudgingError, match=r"ZH-002 \('Senior'\)"):
        read_key(tmp_path / "bad.xlsx")


def test_rerunning_results_keeps_the_owners_choices(comp, tmp_path):
    results = tmp_path / "results.xlsx"
    shutil.copy(comp / "results.xlsx", results)
    demo.choose_winners_and_shortlist(results)
    code, out = run("results", "--scores", str(comp / "scoring.xlsx"), "--key", str(comp / "key.xlsx"), "--out", str(results))
    assert code == 0 and "Kept your Shortlist/Winner choices" in out
    junior = {r["Entry ID"]: r for r in rows(results, "Junior")}
    assert junior["ZH-001"]["Winner (yes/no)"] == "yes" and junior["ZH-002"]["Winner (yes/no)"] == "no"
    assert (tmp_path / "results.previous.xlsx").is_file()


# --------------------------------------------------------------------------- certificates and the scheme

@pytest.fixture(scope="module")
def issued(comp, fonts, tmp_path_factory):
    """The six demo certificates: 2 winners and 4 more shortlisted."""
    results = comp / "results-chosen.xlsx"
    shutil.copy(comp / "results.xlsx", results)
    demo.choose_winners_and_shortlist(results)
    out, work = tmp_path_factory.mktemp("out"), tmp_path_factory.mktemp("work")
    summary = issue_certificates(results, Tokens(REPO), out, None, __import__("datetime").date(2027, 2, 3), work_dir=work)
    return {"out": out, "work": work, "summary": summary}


def test_awards_for_winners_and_shortlisted_in_each_category(issued):
    awards = {s["entry_id"]: s["award"] for s in issued["summary"]}
    assert awards == {"ZH-001": "Winner, Junior", "ZH-002": "Shortlisted, Junior", "ZH-003": "Shortlisted, Junior",
                      "ZH-009": "Winner, Open", "ZH-010": "Shortlisted, Open", "ZH-011": "Shortlisted, Open"}
    assert len(list((issued["out"] / "pdfs").glob("*.pdf"))) == 6      # a winner who is also shortlisted gets one, not two


def test_every_demo_code_opens_with_the_sites_own_decryptCertificate(issued):
    """The records must be scheme v2 exactly: open each one with src/lib/certificateScheme.ts itself."""
    done = subprocess.run(["node", str(TOOL_DIR / "tests" / "decrypt_check.mjs"), str(issued["out"] / "certificates.json"),
                           str(issued["out"] / LEDGER_FILE)], capture_output=True, text=True, cwd=TOOL_DIR)
    assert done.returncode == 0, done.stderr
    opened = json.loads(done.stdout)
    assert opened.pop("_ids") is True                                  # every record id is codeId(code), none extra
    ledger = json.loads((issued["out"] / LEDGER_FILE).read_text())
    assert len(opened) == 6
    for entry in ledger.values():
        assert re.fullmatch(r"[0-9A-HJKMNP-TV-Z]{4}-[0-9A-HJKMNP-TV-Z]{4}", entry["code"])
        assert opened[entry["code"]] == {"name": entry["name"], "award": entry["award"],
                                         "edition": "Inaugural Edition: Zero Hunger", "issued": "2027-02-03"}


def test_the_public_list_has_nothing_readable_and_no_ranking_hint(issued):
    text = (issued["out"] / "certificates.json").read_text()
    records = json.loads(text)
    assert all(sorted(r) == ["data", "id", "iv"] for r in records)    # the build check's rule: exactly id, iv, data
    assert [r["id"] for r in records] == sorted(r["id"] for r in records)
    ledger = json.loads((issued["out"] / LEDGER_FILE).read_text())
    for entry in ledger.values():
        for secret in (entry["name"], entry["code"], entry["code"].replace("-", ""), "Winner", "Shortlisted"):
            assert secret not in text
    assert not list(issued["work"].iterdir())                          # the temporary plain list was deleted


def test_the_site_build_accepts_the_list_and_refuses_readable_data(issued, tmp_path):
    def build(file: Path) -> subprocess.CompletedProcess:
        env = {**__import__("os").environ, "CERTIFICATES_FILE": str(file), "ASTRO_OUT_DIR": str(tmp_path / "dist")}
        return subprocess.run(["npx", "astro", "build"], cwd=REPO, env=env, capture_output=True, text=True)

    good = build(issued["out"] / "certificates.json")
    assert good.returncode == 0, good.stdout + good.stderr
    plain = tmp_path / "plain.json"
    plain.write_text(json.dumps([{"name": "Test Student 01", "award": "Winner, Junior", "code": "ABCD-EFGH"}]))
    bad = build(plain)
    assert bad.returncode != 0 and "unexpected fields" in bad.stdout + bad.stderr


def test_pdf_text(issued):
    pypdf = pytest.importorskip("pypdf")
    ledger = json.loads((issued["out"] / LEDGER_FILE).read_text())
    for entry in ledger.values():
        text = " ".join(pypdf.PdfReader(issued["out"] / "pdfs" / entry["file"]).pages[0].extract_text().split())
        kind, category = entry["award"].split(", ")
        assert kind in text and f"{category} category" in text
        assert entry["name"] in text
        assert statement(Recipient("x", entry["name"], category, kind), EDITION).replace("  ", " ") in text
        assert entry["code"].replace(" ", "") in text.replace(" ", "")
        assert "Check it at pharosprize.org/verify/" in text
        assert "3 February 2027" in text
        assert f"A fundraiser for {SITE['soup_kitchen']}" in text
        for banned in ("United Nations", "Global Goals", "charity", "non-profit", "NGO", "foundation", "society"):
            assert banned.lower() not in text.lower()
        assert not re.search(r"\bUN\b", text)


def test_certificates_keep_their_codes_when_re_run(comp, issued, fonts, tmp_path):
    results = comp / "results-chosen.xlsx"
    before = json.loads((issued["out"] / LEDGER_FILE).read_text())
    again_out = tmp_path / "out"
    shutil.copytree(issued["out"], again_out)
    issue_certificates(results, Tokens(REPO), again_out, None, __import__("datetime").date(2027, 6, 1), work_dir=tmp_path / "w")
    after = json.loads((again_out / LEDGER_FILE).read_text())
    assert {k: (v["code"], v["issued"]) for k, v in after.items()} == {k: (v["code"], v["issued"]) for k, v in before.items()}


def test_nobody_chosen_or_a_typo_stops_certificates(comp, fonts, tmp_path):
    with pytest.raises(JudgingError, match="Nobody is marked"):
        issue_certificates(comp / "results.xlsx", Tokens(REPO), tmp_path / "o", None, __import__("datetime").date(2027, 2, 3), work_dir=tmp_path / "w")
    wb = openpyxl.load_workbook(comp / "results.xlsx")
    ws = wb["Junior"]
    header = [str(c.value or "").lower() for c in ws[1]]
    ws.cell(row=2, column=next(i for i, h in enumerate(header) if h.startswith("winner")) + 1, value="yse")
    wb.save(tmp_path / "typo.xlsx")
    with pytest.raises(JudgingError, match="yse"):
        read_recipients(tmp_path / "typo.xlsx")


def test_write_site_is_the_only_way_to_touch_the_site_file(comp, fonts, tmp_path):
    """A throwaway copy of the site: without --write-site its certificates.json stays [], with it the records land there."""
    site = tmp_path / "site"
    for name in ("tokens.json", "astro.config.mjs", "src/site.ts", "src/lib/certificateScheme.ts", "scripts/encrypt-certificates.mjs"):
        (site / name).parent.mkdir(parents=True, exist_ok=True)
        shutil.copy(REPO / name, site / name)
    (site / "src/data").mkdir(parents=True)
    site_file = site / "src/data/certificates.json"
    site_file.write_text("[]\n")
    results = tmp_path / "results.xlsx"
    shutil.copy(comp / "results.xlsx", results)
    demo.choose_winners_and_shortlist(results)
    args = ["--repo-root", str(site), "certificates", "--results", str(results), "--out", str(tmp_path / "out"), "--issued", "2027-02-03"]
    code, out = run(*args)
    assert code == 0, out
    assert site_file.read_text() == "[]\n" and (tmp_path / "out" / "certificates.json").is_file()
    code, out = run(*args, "--write-site")
    assert code == 0, out
    assert len(json.loads(site_file.read_text())) == 6 and "--write-site" in out
    assert not (REPO / "src/data/certificates.json").read_text().strip().strip("[]").strip()   # the real site file was never touched


# --------------------------------------------------------------------------- privacy and housekeeping

def test_nothing_in_data_is_tracked_by_git_and_data_is_ignored():
    if shutil.which("git") is None or not (REPO / ".git").exists():
        pytest.skip("not a git checkout")
    tracked = subprocess.run(["git", "ls-files", "tools/judging/data", "tools/judging/demo-output", "tools/judging/.fonts"],
                             cwd=REPO, capture_output=True, text=True).stdout.split()
    assert tracked == []
    for path in ("data/key.xlsx", "data/essays/ZH-001.pdf", "data/scoring.xlsx", "data/results.xlsx", "data/out/certificates.json",
                 "data/out/pdfs/x.pdf", "data/out/certificate-codes.json", "demo-output/out/certificates.json", ".fonts/PharosSans.ttf"):
        ignored = subprocess.run(["git", "check-ignore", "-q", f"tools/judging/{path}"], cwd=REPO).returncode
        assert ignored == 0, f"tools/judging/{path} is not git-ignored"


def test_no_ai_calls_and_no_old_scheme_left():
    sources = [p for p in TOOL_DIR.rglob("*") if p.is_file() and p.suffix in (".py", ".mjs", ".md", ".yaml", ".txt")
               and ".venv" not in p.parts and "demo-output" not in p.parts and "tests" not in p.parts]
    assert sources
    ai = re.compile(r"^\s*(?:import|from)\s+(anthropic|openai|google\.generativeai|cohere|mistralai|langchain|ollama)\b", re.M)
    old = re.compile("|".join(["sealing", "verify-reference", "PH-ZH", r"web\.app", r"version 1\b", r"\bv1\b"]), re.I)
    for path in sources:
        text = path.read_text(encoding="utf-8")
        assert not ai.search(text), f"{path.name} imports an AI library"
        assert not old.search(text), f"{path.name} still mentions the old certificate scheme"
    assert not (TOOL_DIR / "sealing.py").exists() and not (TOOL_DIR / "verify-reference.mjs").exists()
    readme = (TOOL_DIR / "README.md").read_text()
    assert "Delete tools/judging/data/ except the published names and certificate records by 3 August 2027" in readme
    assert re.search(r"No AI or LLM", readme)
