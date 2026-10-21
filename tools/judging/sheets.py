"""Make one Excel scoring sheet per judge.

The sheets only ever contain entry IDs, never names: this command does not read key.xlsx.
"""
from __future__ import annotations

import json
import random
from datetime import datetime
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill, Protection
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from common import JudgingError, Tokens, safe_filename
from rubric import Rubric

META_SHEET = "_meta"
SCORES_SHEET = "Scores"
ASSIGNMENTS_FILE = "assignments.json"


def parse_judges(text: str) -> list[str]:
    judges = [j.strip() for j in text.split(",") if j.strip()]
    if len(judges) != len(set(j.lower() for j in judges)):
        raise JudgingError("Each judge name must be different.")
    if len(judges) < 2:
        raise JudgingError("Give at least two judges, e.g. --judges \"A,B,C\".")
    files = [safe_filename(j).lower() for j in judges]
    if len(files) != len(set(files)):
        raise JudgingError("Two judge names would produce the same file name; make them more distinct.")
    return judges


def assign(entry_ids: list[str], judges: list[str], per_essay: int | None,
           seed: int | None) -> dict[str, list[str]]:
    """Decide which entries each judge scores.

    With per_essay None every judge gets every entry. Otherwise each entry goes to
    `per_essay` different judges, always picking the judges with the lightest load so
    workloads differ by at most one essay.
    """
    if per_essay is None or per_essay >= len(judges):
        return {j: sorted(entry_ids) for j in judges}
    if per_essay < 1:
        raise JudgingError("--per-essay must be at least 1 (2 or more lets us compare judges).")
    rng = random.Random(seed)
    load = {j: 0 for j in judges}
    plan: dict[str, list[str]] = {j: [] for j in judges}
    order = list(entry_ids)
    rng.shuffle(order)
    for entry_id in order:
        # Lightest load first; random tie-break so the same judges are not always chosen.
        chosen = sorted(judges, key=lambda j: (load[j], rng.random()))[:per_essay]
        for j in chosen:
            plan[j].append(entry_id)
            load[j] += 1
    return {j: sorted(ids) for j, ids in plan.items()}


def build_sheet(path: Path, judge: str, entry_ids: list[str], rubric: Rubric,
                tokens: Tokens | None) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = SCORES_SHEET

    head_fill = PatternFill("solid", fgColor=(tokens.color_or("cream", "FBEEE8") if tokens else "FBEEE8").lstrip("#"))
    head_font = Font(bold=True, color=(tokens.color_or("plum", "3A2226") if tokens else "3A2226").lstrip("#"))
    first_row, last_row = 2, len(entry_ids) + 1
    n = len(rubric.criteria)
    comments_col = n + 2

    headers = ["Entry ID", *[rubric.header(c) for c in rubric.criteria], "Comments", "Row complete?"]
    for col, text in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col, value=text)
        cell.font, cell.fill = head_font, head_fill
        cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center" if 1 < col <= n + 1 else "left")
    ws.row_dimensions[1].height = 46

    unlocked = Protection(locked=False)
    first_score, last_score = get_column_letter(2), get_column_letter(n + 1)
    for i, entry_id in enumerate(entry_ids):
        row = first_row + i
        ws.cell(row=row, column=1, value=entry_id).font = Font(name="Consolas", bold=True)
        for col in range(2, n + 2):
            cell = ws.cell(row=row, column=col)
            cell.protection = unlocked
            cell.alignment = Alignment(horizontal="center")
        comment = ws.cell(row=row, column=comments_col)
        comment.protection = unlocked
        comment.alignment = Alignment(wrap_text=True, vertical="top")
        ws.cell(row=row, column=comments_col + 1,
                value=f'=IF(COUNT({first_score}{row}:{last_score}{row})={n},"Yes","")').alignment = Alignment(horizontal="center")

    # Only valid scores can be typed into the score cells.
    lo, hi = rubric.scale_min, rubric.scale_max
    kind = "whole" if rubric.integer_only else "decimal"
    validation = DataValidation(
        type=kind, operator="between", formula1=str(lo), formula2=str(hi),
        allow_blank=True, showErrorMessage=True, errorStyle="stop",
        errorTitle="Score not allowed",
        error=f"Enter a {'whole number' if rubric.integer_only else 'number'} from {lo} to {hi}.",
        showInputMessage=True, promptTitle="Score", prompt=f"{lo} (weakest) to {hi} (strongest)",
    )
    validation.add(f"{first_score}{first_row}:{last_score}{last_row}")
    ws.add_data_validation(validation)

    ws.column_dimensions["A"].width = 12
    for col in range(2, n + 2):
        ws.column_dimensions[get_column_letter(col)].width = 18
    ws.column_dimensions[get_column_letter(comments_col)].width = 60
    ws.column_dimensions[get_column_letter(comments_col + 1)].width = 14
    ws.freeze_panes = "B2"

    # Lock everything except score and comment cells, so Entry IDs cannot be changed by accident.
    # No password: this guards against slips, it is not security.
    ws.protection.sheet = True
    ws.protection.formatColumns = False
    ws.protection.formatRows = False

    guide = wb.create_sheet("Rubric")
    guide["A1"] = f"Scoring guide for {judge}"
    guide["A1"].font = Font(bold=True, size=14)
    guide["A2"] = (f"Give every essay a score from {lo} to {hi} for each criterion"
                   f"{' (whole numbers)' if rubric.integer_only else ''}. "
                   "Comments are optional but help the organisers. Leave a row blank if you cannot score it.")
    if rubric.placeholder:
        guide["A3"] = "NOTE: these criteria are PLACEHOLDERS and will be replaced before real judging."
        guide["A3"].font = Font(bold=True, color="C85A55")
    guide.append([])
    guide.append(["Criterion", "Weight", "What to look for"])
    for cell in guide[5]:
        cell.font = Font(bold=True)
    for c in rubric.criteria:
        guide.append([c.name, c.weight, c.description])
    guide.column_dimensions["A"].width = 30
    guide.column_dimensions["B"].width = 10
    guide.column_dimensions["C"].width = 90

    # Hidden sheet that lets `results` read this file reliably, whatever the file is renamed to.
    meta = wb.create_sheet(META_SHEET)
    meta["A1"] = json.dumps({
        "version": 1, "judge": judge, "entries": entry_ids,
        "criteria": [c.id for c in rubric.criteria],
        "created": datetime.now().isoformat(timespec="seconds"),
    })
    meta.sheet_state = "hidden"
    wb.active = 0
    wb.save(path)


def make_sheets(entry_ids: list[str], judges: list[str], rubric: Rubric, out_dir: Path,
                tokens: Tokens | None, per_essay: int | None = None, seed: int | None = None) -> dict[str, list[str]]:
    plan = assign(entry_ids, judges, per_essay, seed)
    out_dir.mkdir(parents=True, exist_ok=True)
    for judge, ids in plan.items():
        build_sheet(out_dir / f"scoring_{safe_filename(judge)}.xlsx", judge, ids, rubric, tokens)
    # Kept so `results` can notice a judge who never sent a sheet back. Holds IDs only.
    (out_dir / ASSIGNMENTS_FILE).write_text(json.dumps(plan, indent=2), encoding="utf-8")
    return plan
