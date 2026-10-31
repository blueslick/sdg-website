"""Read the owner's scoring sheet and write results.xlsx: Junior and Open ranked separately.

The toolkit only adds up marks and sorts by them. It never picks the shortlist or the winners:
those columns (Shortlist yes/no, Winner yes/no) start empty and the owner fills them in himself.
Equal totals share a rank and are shown as ties; nothing breaks them.
Names are joined from key.xlsx here and nowhere else.
"""
from __future__ import annotations

import json
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from common import CATEGORIES, JudgingError, Tokens, normalise_id
from rubric import Rubric
from sheet import (META_SHEET, SCORES_SHEET, SECOND, SECOND_AGAIN, TRIAGE, TRIAGE_MAYBE, TRIAGE_NO,
                   TRIAGE_STRONG, layout)

YES_NO = ("yes", "no")
NOT_RANKED_SHEET = "Not ranked"


@dataclass
class Essay:
    entry_id: str
    triage: str = ""                                   # one of TRIAGE, or "" if not triaged
    marks: dict[str, int] = field(default_factory=dict)
    total: int | None = None                           # only when every criterion has a mark
    comment: str = ""
    second: str = ""                                   # Agree / Read again / ""
    category: str = ""                                 # from key.xlsx
    rank: int | None = None
    reason: str = ""                                   # why it is not ranked
    flags: list[str] = field(default_factory=list)


@dataclass
class Outcome:
    essays: dict[str, Essay]
    flags: list[dict[str, str]]
    kept_picks: int = 0

    def ranked(self, category: str) -> list[Essay]:
        rows = [e for e in self.essays.values() if e.category == category and e.rank is not None]
        return sorted(rows, key=lambda e: (e.rank, e.entry_id))

    def not_ranked(self) -> list[Essay]:
        return sorted((e for e in self.essays.values() if e.rank is None), key=lambda e: e.entry_id)


# --------------------------------------------------------------------------- reading

def _match(value: object, options: tuple[str, ...]) -> str | None:
    text = str(value or "").strip().lower()
    return next((o for o in options if o.lower() == text), None)


def _read_mark(value: object, limit: int) -> tuple[int | None, str | None]:
    """Return (mark, problem). A blank cell is (None, None)."""
    if value is None or (isinstance(value, str) and not value.strip()):
        return None, None
    try:
        number = float(str(value).strip()) if isinstance(value, str) else float(value)
    except (TypeError, ValueError):
        return None, f"'{value}' is not a number"
    if number != int(number):
        return None, f"{number:g} is not a whole number"
    if not 0 <= number <= limit:
        return None, f"{number:g} is outside 0 to {limit}"
    return int(number), None


def read_scores(path: Path, rubric: Rubric) -> tuple[dict[str, Essay], list[dict[str, str]]]:
    if not path.is_file():
        raise JudgingError(f"Scoring sheet not found: {path}. Make it with `sheet`, then fill it in.")
    wb = openpyxl.load_workbook(path, data_only=True)
    if META_SHEET not in wb.sheetnames or SCORES_SHEET not in wb.sheetnames:
        raise JudgingError(f"{path.name} is not a scoring sheet made by this toolkit (run `sheet`).")
    meta = json.loads(wb[META_SHEET]["A1"].value)
    if meta.get("version") != 2:
        raise JudgingError(f"{path.name} was made by an older version of the toolkit. Make a new one with `sheet`.")
    expected = [{"id": c.id, "marks": c.marks} for c in rubric.criteria]
    if meta["criteria"] != expected:
        raise JudgingError(f"{path.name} was made with different criteria or marks from the current rubric.yaml. "
                           "Restore rubric.yaml, or make a new sheet with `sheet` and copy your marks across.")
    col = layout(rubric)
    n = len(rubric.criteria)
    essays: dict[str, Essay] = {}
    flags: list[dict[str, str]] = []

    def flag(kind: str, entry_id: str, detail: str) -> None:
        flags.append({"type": kind, "entry": entry_id, "detail": detail})
        if entry_id in essays:
            essays[entry_id].flags.append(kind)

    seen: set[str] = set()
    for row in wb[SCORES_SHEET].iter_rows(min_row=2, values_only=True):
        entry_id = normalise_id(row[0])
        if not entry_id:
            continue
        if entry_id not in meta["entries"]:
            flag("Unexpected row", entry_id, "Not one of the essays on the sheet; ignored.")
            continue
        seen.add(entry_id)
        cells = list(row) + [None] * (col["second"] - len(row))
        essay = essays.setdefault(entry_id, Essay(entry_id))
        raw_triage = cells[col["triage"] - 1]
        essay.triage = _match(raw_triage, TRIAGE) or ""
        if raw_triage not in (None, "") and not essay.triage:
            flag("Triage not recognised", entry_id, f"'{raw_triage}' is not Clearly no, Maybe or Strong; treated as not triaged.")
        essay.comment = str(cells[col["comments"] - 1] or "").strip()

        entered = [cells[col["first_score"] - 1 + i] for i in range(n)]
        if essay.triage in (TRIAGE_MAYBE, TRIAGE_STRONG):
            for criterion, value in zip(rubric.criteria, entered):
                mark, problem = _read_mark(value, criterion.marks)
                if problem:
                    flag("Invalid mark", entry_id, f"{criterion.name}: {problem}; treated as blank.")
                elif mark is not None:
                    essay.marks[criterion.id] = mark
            if len(essay.marks) == n:
                essay.total = sum(essay.marks.values())
            else:
                missing = [c.name for c in rubric.criteria if c.id not in essay.marks]
                flag("Incomplete marks", entry_id, f"Missing: {', '.join(missing)}. Not ranked until every criterion has a mark.")
        elif any(v not in (None, "") for v in entered):
            flag("Marks ignored", entry_id, f"Marks were typed in, but the triage is "
                 f"{essay.triage or 'blank'}, so they are not used. Set the triage to Maybe or Strong to use them.")

        raw_second = cells[col["second"] - 1]
        if raw_second not in (None, ""):
            second = _match(raw_second, SECOND)
            if not second:
                flag("Second reader not recognised", entry_id, f"'{raw_second}' is not Agree or Read again; ignored.")
            elif essay.triage != TRIAGE_NO:
                flag("Second reader ignored", entry_id, "A second reader is only used on Clearly-no essays; ignored.")
            else:
                essay.second = second
                if second == SECOND_AGAIN:
                    flag("Read again", entry_id, "The second reader says this Clearly-no essay should be read again. "
                         "Change its triage if you agree.")
    for entry_id in meta["entries"]:
        if entry_id not in seen:
            essays.setdefault(entry_id, Essay(entry_id))
            flag("Row deleted", entry_id, "The row was deleted from the scoring sheet.")
    return essays, flags


def read_picks(path: Path) -> dict[str, dict[str, str]]:
    """The owner's Shortlist / Winner choices from an existing results.xlsx: {entry ID: {shortlist, winner}}."""
    picks: dict[str, dict[str, str]] = {}
    if not path.is_file():
        return picks
    wb = openpyxl.load_workbook(path, data_only=True)
    for category in CATEGORIES:
        if category not in wb.sheetnames:
            continue
        rows = wb[category].iter_rows(values_only=True)
        header = [str(h or "").strip().lower() for h in next(rows, [])]
        cols = pick_columns(header)
        if cols is None:
            continue
        for row in rows:
            entry_id = normalise_id(row[cols["id"]]) if len(row) > cols["id"] else ""
            if entry_id:
                picks[entry_id] = {"shortlist": str(row[cols["shortlist"]] or "").strip(),
                                   "winner": str(row[cols["winner"]] or "").strip()}
    return picks


def pick_columns(header: list[str]) -> dict[str, int] | None:
    """Positions of the columns the owner fills in, found by their header text. None if not a results sheet."""
    def find(prefix: str) -> int | None:
        return next((i for i, h in enumerate(header) if h.startswith(prefix)), None)
    found = {"id": find("entry id"), "name": find("name"), "shortlist": find("shortlist"), "winner": find("winner")}
    return None if any(v is None for v in found.values()) else found  # type: ignore[return-value]


# --------------------------------------------------------------------------- computing

def compute(scores: Path, rubric: Rubric, key: dict[str, dict[str, str]], essay_ids: set[str] | None = None) -> Outcome:
    essays, flags = read_scores(scores, rubric)

    def flag(kind: str, entry_id: str, detail: str) -> None:
        flags.append({"type": kind, "entry": entry_id, "detail": detail})
        if entry_id in essays:
            essays[entry_id].flags.append(kind)

    for entry_id in sorted((essay_ids or set()) - essays.keys()):
        essays[entry_id] = Essay(entry_id, reason="Not on the scoring sheet")
        flag("Not on the scoring sheet", entry_id, "The essay file exists but is not on the scoring sheet. Make a new sheet.")
    for entry_id in sorted(key.keys() - essays.keys()):
        if essay_ids is not None and entry_id not in essay_ids:
            flags.append({"type": "No essay file", "entry": entry_id, "detail": "Listed in key.xlsx but there is no essay file."})
            continue
        essays[entry_id] = Essay(entry_id, reason="Not on the scoring sheet")
        flag("Not on the scoring sheet", entry_id, "In key.xlsx but not on the scoring sheet.")
    for entry_id, essay in essays.items():
        who = key.get(entry_id)
        if who:
            essay.category = who["category"]
        else:
            flag("Not in key", entry_id, "No row in key.xlsx, so there is no name or category. It cannot be ranked.")

    # Reasons an essay is not ranked, then ranks within each category.
    for essay in essays.values():
        if essay.reason:
            continue
        if not essay.category:
            essay.reason = "Not in key.xlsx (no category)"
        elif essay.triage == TRIAGE_NO:
            essay.reason = "Clearly no"
        elif not essay.triage:
            essay.reason = "Not triaged yet"
            flag("Not triaged", essay.entry_id, "No triage has been chosen for this essay.")
        elif essay.total is None:
            essay.reason = "Marks missing"

    for category in CATEGORIES:
        pool = sorted((e for e in essays.values() if e.category == category and not e.reason and e.total is not None),
                      key=lambda e: (-e.total, e.entry_id))
        for position, essay in enumerate(pool, start=1):
            before = pool[position - 2] if position > 1 else None
            essay.rank = before.rank if before and before.total == essay.total else position
        by_rank: dict[int, list[Essay]] = defaultdict(list)
        for essay in pool:
            by_rank[essay.rank].append(essay)
        for rank, group in by_rank.items():
            if len(group) > 1:
                ids = ", ".join(e.entry_id for e in group)
                for e in group:
                    flag("Tie", e.entry_id, f"{category}: tied on {e.total} at rank {rank} with {ids}. Not broken: you decide.")
    return Outcome(essays, flags)


# --------------------------------------------------------------------------- writing

def write_results(outcome: Outcome, rubric: Rubric, key: dict[str, dict[str, str]], path: Path,
                  tokens: Tokens | None, picks: dict[str, dict[str, str]] | None = None) -> None:
    def colour(name: str, fallback: str) -> str:
        return (tokens.color_or(name, fallback) if tokens else fallback).lstrip("#")

    head_fill = PatternFill("solid", fgColor=colour("cream", "FBEEE8"))
    head_font = Font(bold=True, color=colour("plum", "3A2226"))
    warn_fill = PatternFill("solid", fgColor=colour("peach", "F3C6AA"))
    picks = picks or {}

    def style_header(ws, widths: list[int]) -> None:
        for cell in ws[1]:
            cell.font, cell.fill = head_font, head_fill
            cell.alignment = Alignment(wrap_text=True, vertical="center")
        for i, width in enumerate(widths, start=1):
            ws.column_dimensions[get_column_letter(i)].width = width
        ws.freeze_panes = "A2"

    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    n = len(rubric.criteria)
    for category in CATEGORIES:
        ws = wb.create_sheet(category)
        ws.append(["Rank", "Entry ID", "Name", "School", "Country", "Triage", *[c.name for c in rubric.criteria],
                   f"Total (out of {rubric.total})", "Shortlist (yes/no)", "Winner (yes/no)", "Your comments", "Flags"])
        rows = outcome.ranked(category)
        for e in rows:
            who = key.get(e.entry_id, {})
            kept = picks.get(e.entry_id, {})
            ws.append([e.rank, e.entry_id, who.get("name", ""), who.get("school", ""), who.get("country", ""), e.triage,
                       *[e.marks.get(c.id) for c in rubric.criteria], e.total,
                       kept.get("shortlist") or None, kept.get("winner") or None, e.comment,
                       "; ".join(dict.fromkeys(e.flags))])
            if e.flags:
                for cell in ws[ws.max_row]:
                    cell.fill = warn_fill
        if rows:
            letters = [get_column_letter(7 + n + 1), get_column_letter(7 + n + 2)]   # Shortlist, Winner
            rule = DataValidation(type="list", formula1='"yes,no"', allow_blank=True, showErrorMessage=True,
                                  errorStyle="stop", errorTitle="yes or no", error="Type yes or no, or leave it empty.")
            for letter in letters:
                rule.add(f"{letter}2:{letter}{len(rows) + 1}")
            ws.add_data_validation(rule)
        style_header(ws, [7, 11, 28, 26, 14, 12, *[16] * n, 14, 14, 12, 40, 28])

    ws = wb.create_sheet(NOT_RANKED_SHEET)
    ws.append(["Entry ID", "Name", "Category", "Triage", "Second reader", "Why it is not ranked", "Your comments"])
    for e in outcome.not_ranked():
        who = key.get(e.entry_id, {})
        ws.append([e.entry_id, who.get("name", ""), e.category, e.triage, e.second, e.reason, e.comment])
        if e.second == SECOND_AGAIN or e.reason not in ("Clearly no",):
            for cell in ws[ws.max_row]:
                cell.fill = warn_fill
    style_header(ws, [11, 28, 11, 12, 14, 34, 50])

    fs = wb.create_sheet("Flags")
    fs.append(["Type", "Entry ID", "Detail"])
    for f in sorted(outcome.flags, key=lambda f: (f["type"], f["entry"])):
        fs.append([f["type"], f["entry"], f["detail"]])
    if not outcome.flags:
        fs.append(["None", "", "Nothing to review."])
    style_header(fs, [26, 11, 110])

    ns = wb.create_sheet("Notes")
    lines = [
        ("CONFIDENTIAL", "Contains names from key.xlsx. Keep it in tools/judging/data/ and do not share it."),
        ("Generated", datetime.now().strftime("%Y-%m-%d %H:%M")),
        ("Rubric", rubric.name + ("  [PLACEHOLDER RUBRIC: DRAFT RESULTS]" if rubric.placeholder else "")),
        ("Total", f"The sum of your marks for the {n} criteria, out of {rubric.total}. Only Maybe and Strong essays with every mark filled in are ranked."),
        ("Rank", "Highest total = 1, separately for Junior and Open. Equal totals share a rank (1, 2, 2, 4) and are flagged as ties. The tool does not break ties; the order inside a tie is only by Entry ID."),
        ("Shortlist, Winner", "Empty on purpose. Type yes or no yourself. The toolkit never chooses the shortlist or the winners. "
                              "`certificates` reads these two columns (a Winner gets a Winner certificate, any other Shortlist = yes gets a Shortlisted one)."),
        ("Name", "`certificates` prints the Name column of the Junior and Open sheets on the certificate. Correct a spelling here if you need to."),
        ("Re-running", "If you run `results` again it keeps the yes/no you already typed for each Entry ID (a copy of the old file is saved as results.previous.xlsx)."),
        ("Not ranked", "Clearly-no essays, essays not yet triaged, and Maybe/Strong essays with missing marks are listed on the Not ranked sheet. "
                       "Clearly-no essays whose second reader said Read again are highlighted."),
        ("Values", "Numbers are calculated by the toolkit and written as values, not live formulas."),
    ]
    for label, text in lines:
        ns.append([label, text])
    for row in ns.iter_rows():
        row[0].font = Font(bold=True)
    ns.column_dimensions["A"].width = 18
    ns.column_dimensions["B"].width = 120
    wb.save(path)
