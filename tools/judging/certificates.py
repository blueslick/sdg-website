"""PDF certificates in the Pharos look, plus the encrypted certificates.json the site reads.

One certificate for each winner ("Winner, Junior" / "Winner, Open") and each other shortlisted
entrant ("Shortlisted, Junior" / "Shortlisted, Open"), exactly as the owner marked them yes in
results.xlsx. The toolkit never decides who gets one.

There is no second copy of the encryption here. The site's own script does it:
    node scripts/encrypt-certificates.mjs <plain> <encrypted> --generate-codes --plain-out <with-codes>
(scheme version 2, defined in src/lib/certificateScheme.ts and documented in src/data/README.md).
This module writes the plain list to a temporary file, runs that script, reads the codes back for the
PDFs, and deletes the temporary file.
"""
from __future__ import annotations

import json
import shutil
import subprocess
import tempfile
import unicodedata
from dataclasses import dataclass
from datetime import date
from pathlib import Path

import openpyxl
from reportlab.lib.colors import HexColor
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.units import mm
from reportlab.lib.utils import simpleSplit
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfgen.canvas import Canvas

from common import CATEGORIES, DATA_DIR, JudgingError, Tokens, safe_filename, site_values
from fontkit import register_fonts
from results import pick_columns

EDITION = "Inaugural Edition: Zero Hunger"
LEDGER_FILE = "certificate-codes.json"       # PRIVATE: names and codes, so a re-run keeps the same codes
SCRIPT = Path("scripts") / "encrypt-certificates.mjs"
WINNER, SHORTLISTED = "Winner", "Shortlisted"


@dataclass(frozen=True)
class Recipient:
    entry_id: str
    name: str
    category: str      # Junior / Open
    kind: str          # Winner / Shortlisted

    @property
    def award(self) -> str:
        return f"{self.kind}, {self.category}"


def statement(recipient: Recipient, edition: str) -> str:
    what = "the winning essay" if recipient.kind == WINNER else "a shortlisted essay"
    return f"for {what} in the {recipient.category} category of the {edition}."


# --------------------------------------------------------------------------- who gets one

def _yes_no(value: object, where: str, column: str) -> bool:
    text = str(value or "").strip().lower()
    if text in ("", "no"):
        return False
    if text == "yes":
        return True
    raise JudgingError(f"results.xlsx, {where}: {column} says '{value}'. Use yes, no, or leave it empty.")


def read_recipients(results_path: Path) -> list[Recipient]:
    """Read the owner's Shortlist and Winner choices (and the names) from results.xlsx."""
    if not results_path.is_file():
        raise JudgingError(f"results.xlsx not found: {results_path}. Run `results` first.")
    wb = openpyxl.load_workbook(results_path, data_only=True)
    out: list[Recipient] = []
    for category in CATEGORIES:
        if category not in wb.sheetnames:
            raise JudgingError(f"results.xlsx has no '{category}' sheet. Run `results` to make it.")
        rows = wb[category].iter_rows(values_only=True)
        cols = pick_columns([str(h or "").strip().lower() for h in next(rows, [])])
        if cols is None:
            raise JudgingError(f"The '{category}' sheet needs Entry ID, Name, Shortlist and Winner columns.")
        for row in rows:
            entry_id = str(row[cols["id"]] or "").strip()
            if not entry_id:
                continue
            where = f"{category} sheet, {entry_id}"
            winner = _yes_no(row[cols["winner"]], where, "Winner")
            shortlisted = _yes_no(row[cols["shortlist"]], where, "Shortlist")
            if not (winner or shortlisted):
                continue
            name = unicodedata.normalize("NFC", str(row[cols["name"]] or "").strip())
            if not name:
                raise JudgingError(f"No name for {entry_id} in results.xlsx. Check key.xlsx has this entry, then re-run `results`.")
            out.append(Recipient(entry_id, name, category, WINNER if winner else SHORTLISTED))
    return out


# --------------------------------------------------------------------------- the site's encrypt script

def run_encrypt_script(repo_root: Path, plain: Path, encrypted: Path, with_codes: Path) -> None:
    script = repo_root / SCRIPT
    if not script.is_file():
        raise JudgingError(f"{SCRIPT} not found in {repo_root}. The toolkit needs the site folder (pass --repo-root).")
    try:
        done = subprocess.run(["node", str(script), str(plain), str(encrypted), "--generate-codes", "--plain-out", str(with_codes)],
                              cwd=repo_root, capture_output=True, text=True)
    except FileNotFoundError:
        raise JudgingError("Node.js was not found. Install it (the site needs it too) and try again.") from None
    if done.returncode != 0:
        raise JudgingError("The site's encrypt script failed:\n" + (done.stderr or done.stdout).strip())


# --------------------------------------------------------------------------- drawing

def _spaced(c: Canvas, text: str, x: float, y: float, font: str, size: float, space: float) -> None:
    t = c.beginText(x, y)
    t.setFont(font, size)
    t.setCharSpace(space)
    t.textOut(text)
    t.setCharSpace(0)          # character spacing otherwise leaks into every later string on the page
    c.drawText(t)


def _fit(text: str, font: str, size: float, min_size: float, max_width: float) -> float:
    while size > min_size and pdfmetrics.stringWidth(text, font, size) > max_width:
        size -= 0.5
    return size


def _slab_path(c: Canvas, width: float, bottom: float, top: float, radius: float):
    """The plum slab: square on three sides, one big round corner bottom-right (as on the site)."""
    k = 0.5523 * radius
    p = c.beginPath()
    p.moveTo(0, bottom)
    p.lineTo(width - radius, bottom)
    p.curveTo(width - radius + k, bottom, width, bottom + radius - k, width, bottom + radius)
    p.lineTo(width, top)
    p.lineTo(0, top)
    p.close()
    return p


def draw_certificate(path: Path, tokens: Tokens, site: dict, recipient: Recipient, edition: str,
                     code: str, issued: date) -> None:
    """One A4-landscape certificate in the site's look: peach ground, plum slab, Gloock display type,
    Instrument Sans body, the Kaushan Script wordmark, rings and the one edition disc. No SDG icons,
    no signatures, no names of judges."""
    W, H = landscape(A4)
    X0 = 22 * mm
    col = {n: HexColor(tokens.color(n)) for n in ("peach", "plum", "plum-muted", "cream", "pink", "coral", "edition")}
    c = Canvas(str(path), pagesize=(W, H))
    c.setTitle(f"{recipient.award} certificate: {recipient.name}")
    c.setAuthor(site["site_name"])
    c.setSubject(f"{edition}. Verification code {code}")

    c.setFillColor(col["peach"]); c.rect(0, 0, W, H, stroke=0, fill=1)
    slab_w, slab_bottom, radius = 255 * mm, H - 100 * mm, 58 * mm
    slab = _slab_path(c, slab_w, slab_bottom, H, radius)
    c.setFillColor(col["plum"]); c.drawPath(slab, stroke=0, fill=1)
    c.saveState()
    c.clipPath(slab, stroke=0, fill=0)
    c.setStrokeColor(col["pink"]); c.setLineWidth(0.5)
    for i in range(1, 11):
        c.circle(slab_w - 6 * mm, slab_bottom + 6 * mm, i * 3.6 * mm, stroke=1, fill=0)
    c.restoreState()
    c.setFillColor(col["edition"]); c.circle(W - 24 * mm, slab_bottom - 4 * mm, 21 * mm, stroke=0, fill=1)   # the one edition shape

    # Wordmark: a decorative accent, so the plain name is also in the PDF metadata.
    c.setFillColor(col["peach"]); c.setFont("PharosScript", 32)
    c.drawString(X0, H - 25 * mm, site["site_name"])

    # Hook: the award in big display type. The same words are plain text in the sentence below.
    c.setFillColor(col["peach"]); c.setFont("PharosSansSemi", 13)
    c.drawString(X0, H - 43 * mm, edition)
    size = _fit(recipient.kind, "PharosDisplay", 78, 40, 185 * mm)
    c.setFillColor(col["cream"]); c.setFont("PharosDisplay", size)
    c.drawString(X0, H - 70 * mm, recipient.kind)
    c.setFillColor(col["cream"]); c.setFont("PharosSansSemi", 18)
    c.drawString(X0, H - 80 * mm, f"{recipient.category} category")
    c.setFillColor(col["coral"]); c.rect(X0, H - 88 * mm, 34 * mm, 2.2 * mm, stroke=0, fill=1)

    # Facts on peach: who, for what, when, and how to check.
    c.setFillColor(col["plum-muted"]); c.setFont("PharosSansSemi", 11)
    c.drawString(X0, 81 * mm, "This certificate is awarded to")
    name_w = 215 * mm
    size = _fit(recipient.name, "PharosDisplay", 42, 16, name_w)
    c.setFillColor(col["plum"]); c.setFont("PharosDisplay", size)
    c.drawString(X0, 66 * mm, recipient.name)
    c.setStrokeColor(col["plum"]); c.setLineWidth(1.4)
    c.line(X0, 60 * mm, X0 + name_w, 60 * mm)

    c.setFont("PharosSans", 14)
    y = 50 * mm
    for line in simpleSplit(statement(recipient, edition), "PharosSans", 14, name_w):
        c.drawString(X0, y, line)
        y -= 19

    c.setFillColor(col["plum-muted"]); c.setFont("PharosSansSemi", 10)
    c.drawString(X0, 33 * mm, "Verification code")
    c.drawString(X0 + 85 * mm, 33 * mm, "Issued")
    c.setFillColor(col["plum"])
    _spaced(c, code, X0, 25 * mm, "PharosSansSemi", 18, 1.0)
    c.setFont("PharosSans", 14)
    c.drawString(X0 + 85 * mm, 25.6 * mm, f"{issued.day} {issued:%B %Y}")
    c.setFont("PharosSans", 11)
    c.drawString(X0, 17 * mm, f"Check it at {site['site_url'].split('://', 1)[-1]}/verify/")

    c.setFillColor(col["plum-muted"]); c.setFont("PharosSans", 9)
    c.drawString(X0, 9 * mm, f"Run by student volunteers. A fundraiser for {site['soup_kitchen']}.")
    c.showPage()
    c.save()


# --------------------------------------------------------------------------- the command

def issue_certificates(results_path: Path, tokens: Tokens, out_dir: Path, site_json: Path | None, issued: date,
                       edition: str = EDITION, allow_missing_glyphs: bool = False,
                       work_dir: Path = DATA_DIR) -> list[dict]:
    """Make the PDFs and the encrypted list. `site_json` is where the list goes (src/data/certificates.json
    with --write-site); otherwise it goes to out_dir/certificates.json."""
    site = site_values(tokens.root)
    recipients = read_recipients(results_path)
    if not recipients:
        raise JudgingError("Nobody is marked Shortlist = yes or Winner = yes in results.xlsx yet, so there is nothing to make. "
                           "Fill in those two columns yourself first.")
    coverage = register_fonts()
    unsupported = {r.entry_id: sorted({ch for ch in r.name if ord(ch) not in coverage["PharosDisplay"]}) for r in recipients}
    unsupported = {k: v for k, v in unsupported.items() if v}
    if unsupported and not allow_missing_glyphs:
        detail = "; ".join(f"{k}: {' '.join(v)}" for k, v in unsupported.items())
        raise JudgingError("These names contain characters the certificate font cannot show (" + detail + "). "
                           "Nothing was written. Use the English spelling in the Name column of results.xlsx, "
                           "or add --allow-missing-glyphs to print them anyway (gaps will appear).")

    out_dir.mkdir(parents=True, exist_ok=True)
    ledger_path = out_dir / LEDGER_FILE
    ledger: dict[str, dict] = json.loads(ledger_path.read_text(encoding="utf-8")) if ledger_path.is_file() else {}

    # Entries whose certificate has not changed keep their code and date, so certificates already sent stay valid.
    plain: list[dict] = []
    for r in recipients:
        old = ledger.get(r.entry_id)
        same = bool(old and (old["name"], old["award"], old["edition"]) == (r.name, r.award, edition))
        if old and not same:
            print(f"Note: {r.entry_id} changed ({old['award']} -> {r.award}, or the name); a new code is issued, "
                  f"and the old code {old['code']} stops working.")
        entry = {"entry_id": r.entry_id, "name": r.name, "award": r.award, "edition": edition,
                 "issued": old["issued"] if same else issued.isoformat()}
        if same:
            entry["code"] = old["code"]
        plain.append(entry)

    work_dir.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory(dir=work_dir, prefix=".tmp-certificates-") as tmp:     # deleted when this block ends
        tmp_dir = Path(tmp)
        plain_file, encrypted_file, with_codes_file = tmp_dir / "plain.json", tmp_dir / "encrypted.json", tmp_dir / "with-codes.json"
        plain_file.write_text(json.dumps(plain, indent=2, ensure_ascii=False), encoding="utf-8")
        run_encrypt_script(tokens.root, plain_file, encrypted_file, with_codes_file)
        issued_list = json.loads(with_codes_file.read_text(encoding="utf-8"))
        records = json.loads(encrypted_file.read_text(encoding="utf-8"))
    if len(issued_list) != len(recipients) or len(records) != len(recipients):
        raise JudgingError("The encrypt script returned a different number of certificates than expected. Nothing was written.")

    # The list is public, so its order must not hint at who placed where: sort by the (hashed) id.
    records.sort(key=lambda rec: rec["id"])
    destination = site_json or out_dir / "certificates.json"
    destination.parent.mkdir(parents=True, exist_ok=True)
    destination.write_text(json.dumps(records, indent=2) + "\n", encoding="utf-8")

    pdf_dir = out_dir / "pdfs"
    shutil.rmtree(pdf_dir, ignore_errors=True)
    pdf_dir.mkdir(parents=True)
    new_ledger: dict[str, dict] = {}
    summary: list[dict] = []
    for r, entry in zip(recipients, issued_list):
        file_name = f"{r.entry_id}_{safe_filename(r.award)}_{entry['code']}.pdf"
        draw_certificate(pdf_dir / file_name, tokens, site, r, edition, entry["code"], date.fromisoformat(entry["issued"]))
        new_ledger[r.entry_id] = {k: entry[k] for k in ("code", "name", "award", "edition", "issued")} | {"file": file_name}
        summary.append({"entry_id": r.entry_id, "name": r.name, "award": r.award, "code": entry["code"], "file": file_name})
    ledger_path.write_text(json.dumps(new_ledger, indent=2, ensure_ascii=False), encoding="utf-8")
    return summary
